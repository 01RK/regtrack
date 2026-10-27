from testdb import new_database
"""Meeting import integration tests against the Access schema and API."""
import copy
from itertools import permutations
import json
import io
from pathlib import Path
import tempfile
import unittest
from app import create_app
from testdb import initialize as init_db
from dbtest import Database
from meeting_import import candidates
from meeting_workbook import SHEETS, read_workbook
from openpyxl import load_workbook
from werkzeug.datastructures import FileStorage

ROOT = Path(__file__).resolve().parents[1]


class MeetingImportTest(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.path = new_database()
        init_db(self.path, with_seed=False)
        self.client = create_app(self.path).test_client()
        self.packet = {'meeting': {'title': '动力蓄电池性能工作组第二次会议', 'meeting_date': '2026-07-15',
                                  'key_discussions': '循环寿命与性能评价', 'organizer': '测试组织'},
                       'standards': [{'key': 'S1', 'name_cn': '电动汽车动力蓄电池循环寿命要求', 'std_no': 'GB/T 31484', 'note': '讨论循环倍率。', 'current_stage': '起草'},
                                     {'key': 'S2', 'name_cn': '电动汽车动力蓄电池性能评价指南', 'std_no': '', 'note': '尚无正式标准号。', 'current_stage': '预研'}],
                       'actions': [{'title': '更新草案', 'description': '按会议意见修订', 'standard_key': 'S1', 'due_date': '2026-09-25'},
                                   {'title': '研究评分', 'description': '研究评分细则', 'standard_key': 'S2'}]}
        self.decision = {'meeting_id': 'new', 'standards': {'S1': {'mode': 'create', 'profile': {'name_cn': self.packet['standards'][0]['name_cn'], 'std_no': 'GB/T 31484'}}, 'S2': {'mode': 'skip'}}, 'confirmed': True}

    def tearDown(self):
        self.tmp.cleanup()

    def rows(self, table):
        with Database(self.path) as conn:
            order = ' ORDER BY id' if 'id' in conn.columns(table) else ''
            return conn.rows(f'SELECT * FROM {table}{order}')

    def preview(self):
        return self.client.post('/api/meetings/imports/preview', data={'file': (io.BytesIO(self.excel()), 'meeting.xlsx')})

    def excel(self):
        workbook = load_workbook(ROOT / 'static' / 'meeting-import-template.xlsx')
        for name, (_, fields) in SHEETS.items():
            records = [self.packet['meeting']] if name == 'Meeting' else self.packet[name.lower()]
            for i, record in enumerate(records, 3):
                for j, field in enumerate(fields, 1):
                    workbook[name].cell(i, j, record.get(field, ''))
        stream = io.BytesIO()
        workbook.save(stream)
        workbook.close()
        return stream.getvalue()

    def review(self):
        return self.client.post('/api/meetings/imports/review', json={'package': self.packet, 'decision': self.decision})

    def commit(self, actions=None):
        if actions is None:
            reviewed = self.review()
            if reviewed.status_code != 200:
                return reviewed
            r = reviewed.json
            actions = {'existing_snapshot': r['existing_snapshot'],
                       'keep_existing': [int(a['key'].split(':')[1]) for a in r['rows'] if a['origin']=='existing' and a['keep']],
                       'keep_new': [int(a['key'].split(':')[1]) for a in r['rows'] if a['origin']=='new' and a['keep']]}
        self.decision['actions'] = actions
        return self.client.post('/api/meetings/imports', json={'package': self.packet, 'decision': self.decision}, headers={'X-User': 'Tester'})

    def test_stakeholder_template_preview_and_persistence(self):
        self.packet['actions'][0]['stakeholder'] = '技术响应单位'
        preview = self.preview()
        self.assertEqual(preview.status_code, 200, preview.json)
        self.assertEqual(preview.json['package']['actions'][0]['stakeholder'], '技术响应单位')
        self.assertNotIn('owner', preview.json['package']['actions'][0])
        self.assertEqual(self.commit().status_code, 201)
        self.assertEqual(self.rows('action_item')[0]['coordinator'], '技术响应单位')
        self.decision['meeting_id'] = self.rows('meeting')[0]['id']
        review = self.review()
        self.assertEqual(review.status_code, 200, review.json)
        existing = next(a for a in review.json['rows'] if a['origin'] == 'existing')
        self.assertEqual(existing['stakeholder'], '技术响应单位')

    def test_read_only_preview_and_assets(self):
        self.assertEqual(self.preview().status_code, 200)
        self.assertEqual(self.rows('meeting'), [])
        for path in ['/meetings/import', '/static/meeting-import-template.xlsx', '/static/meeting-import-prompt.txt']:
            with self.client.get(path) as response:
                self.assertEqual(response.status_code, 200)

    def test_create_tracked_reference_actions_and_history(self):
        result = self.commit()
        self.assertEqual(result.status_code, 201, result.json)
        self.assertEqual(len(self.rows('standard')), 1)
        self.assertEqual(len(self.rows('standard_stage_history')), 1)
        self.assertEqual(len(self.rows('action_status_history')), 1)
        detail = self.client.get(f'/api/meetings/{result.json["meeting_id"]}').json
        self.assertNotIn('references', detail)
        self.assertEqual(len(detail['standards']), 1)
        self.assertEqual(len(detail['actions']), 1)
        actions = self.rows('action_item')
        self.assertIsNotNone(actions[0]['standard_id'])
        self.assertEqual(result.json['excluded_actions'], 1)
        self.assertFalse(any(a['title'] == '研究评分' for a in actions))
        self.assertEqual(actions[0]['created_by'], 'Tester')

    def test_reimport_is_idempotent_and_preserves_manual_notes(self):
        first = self.commit().json
        sid = self.rows('standard')[0]['id']
        self.decision['meeting_id'] = first['meeting_id']
        self.decision['standards']['S1'] = {'mode': 'existing', 'id': sid}
        self.client.put(f'/api/meetings/{first["meeting_id"]}/standards/{sid}', json={'note': '人工批注'})
        for _ in range(2):
            response = self.commit()
            self.assertEqual(response.status_code, 201, response.json)
            self.assertEqual(response.json['omitted_actions'], 1)
        self.assertEqual(len(self.rows('action_item')), 1)
        self.assertEqual(self.rows('meeting_standard')[0]['note'], '人工批注\n\n讨论循环倍率。')

    def test_invalid_decision_rolls_back_everything(self):
        self.decision['standards']['S2'] = {'mode': 'existing', 'id': 999}
        result = self.commit()
        self.assertEqual(result.status_code, 400)
        for table in ['meeting', 'standard', 'meeting_standard', 'lookup_value', 'standard_stage_history', 'action_item']:
            self.assertEqual(self.rows(table), [], table)

    def test_existing_match_found_by_search_outside_candidates(self):
        """没有推荐结果时，用户可搜索任一已有标准作为匹配对象。"""
        other = self.client.post('/api/standards', json={'std_no': 'GB/T 99001-2026', 'name_cn': '完全不同的名称',
                                                         'stage_code': 'PRE_RESEARCH'}, headers={'X-User': 'Tester'}).json
        preview = self.preview().json
        self.assertNotIn(other['id'], [c['id'] for match in preview['standards'] for c in match['candidates']])
        self.decision['standards']['S2'] = {'mode': 'existing', 'id': other['id']}
        result = self.commit()
        self.assertEqual(result.status_code, 201, result.json)
        meeting = self.client.get(f"/api/meetings/{result.json['meeting_id']}").json
        self.assertIn(other['id'], [s['id'] for s in meeting['standards']])
        self.assertEqual(self.client.get(f"/api/standards/{other['id']}").json['stage_code'], 'PRE_RESEARCH')

    def test_two_minutes_cannot_match_the_same_searched_standard(self):
        other = self.client.post('/api/standards', json={'name_cn': '被重复匹配的标准', 'stage_code': 'DRAFTING'},
                                 headers={'X-User': 'Tester'}).json
        self.decision['standards'] = {'S1': {'mode': 'existing', 'id': other['id']},
                                      'S2': {'mode': 'existing', 'id': other['id']}}
        self.assertEqual(self.commit().status_code, 400)
        self.assertEqual(self.rows('meeting'), [])

    def test_tracking_standards_without_numbers(self):
        self.decision['standards']['S1']['profile']['std_no'] = ''
        self.decision['standards']['S2'] = {'mode': 'create', 'profile': {'name_cn': '性能评价指南', 'std_no': ''}}
        result = self.commit()
        self.assertEqual(result.status_code, 201, result.json)
        self.assertEqual(result.json['created_standards'], 2)
        self.assertTrue(all(s['std_no'] is None for s in self.rows('standard')))
        self.assertEqual(len(self.rows('meeting_standard')), 2)

    def test_multiple_standards_can_share_provisional_number(self):
        self.decision['standards']['S1']['profile']['std_no'] = 'GB/T XXXX'
        self.decision['standards']['S2'] = {'mode': 'create', 'profile': {'name_cn': '性能评价指南', 'std_no': 'GB/T XXXX'}}
        result = self.commit()
        self.assertEqual(result.status_code, 201, result.json)
        self.assertEqual([s['std_no'] for s in self.rows('standard')], ['GB/T XXXX', 'GB/T XXXX'])

    def test_duplicate_formal_number_rolls_back(self):
        self.decision['standards']['S2'] = {'mode': 'create', 'profile': {'name_cn': '性能评价指南', 'std_no': 'GB/T 31484'}}
        self.assertEqual(self.commit().status_code, 400)
        self.assertEqual(self.rows('standard'), [])
        self.assertEqual(self.rows('meeting'), [])

    def test_standard_profile_accepts_empty_number(self):
        for name in ['待立项甲', '待立项乙']:
            result = self.client.post('/api/standards', json={'name_cn': name, 'std_no': '', 'stage_code': 'PRE_RESEARCH'})
            self.assertEqual(result.status_code, 201, result.json)
            self.assertIsNone(result.json['std_no'])

    def test_exchange_and_lifecycle_with_unassigned_numbers(self):
        self.decision['standards']['S1']['profile']['std_no'] = 'GB/T XXXX'
        self.decision['standards']['S2'] = {'mode': 'create', 'profile': {'name_cn': '性能评价指南', 'std_no': 'GB/T XXXX'}}
        for key in ['S3', 'S4']:
            self.packet['standards'].append({'key': key, 'name_cn': f'未编号标准{key}', 'std_no': '', 'note': '讨论测试方法。', 'current_stage': '起草'})
            self.decision['standards'][key] = {'mode': 'create', 'profile': {'name_cn': f'未编号标准{key}', 'std_no': ''}}
        self.assertEqual(self.commit().status_code, 201)
        sid = next(s['id'] for s in self.rows('standard') if s['std_no'] is None)
        with self.client.get(f'/api/standards/{sid}/lifecycle.md') as response:
            self.assertEqual(response.status_code, 200)
            self.assertNotIn('None', response.data.decode('utf-8'))
        with self.client.get('/api/transfer/export?user=Tester') as response:
            bundle = response.data
        path = new_database()
        init_db(path, with_seed=False)
        other = create_app(path).test_client()
        for _ in range(2):
            result = other.post('/api/transfer/import', data={'file': (io.BytesIO(bundle), 'exchange.json'), 'as_user': 'Receiver'})
            self.assertEqual(result.status_code, 200, result.json)
            standards = other.get('/api/standards').json['items']
            self.assertEqual(len(standards), 4)
            self.assertEqual(sum(s['std_no'] is None for s in standards), 2)
            self.assertEqual(sum(s['std_no'] == 'GB/T XXXX' for s in standards), 2)

    def test_matching_threshold_limit_and_date(self):
        title = self.packet['meeting']['title']
        for i in range(5):
            self.client.post('/api/meetings', json={'title': title + str(i), 'meeting_date': '2026-07-15'})
        self.client.post('/api/meetings', json={'title': title, 'meeting_date': '2026-07-16'})
        result = self.preview().json['meetings']
        self.assertEqual(len(result), 6)
        self.assertEqual(sum(m['same_day'] for m in result), 5)
        self.assertEqual(sum(m['similar_title'] for m in result), 3)
        self.assertEqual(len({m['id'] for m in result}), 6)
        rows = [{'id': i, 'name_cn': title + str(i)} for i in range(5)] + [{'id': 9, 'name_cn': '无关内容'}]
        self.assertEqual(len(candidates(title, rows, 'name_cn')), 3)
        self.assertEqual(candidates('abcdefg', [{'id': 1, 'name_cn': 'abcxxxx'}], 'name_cn'), [])

    def test_existing_meeting_keeps_identity_and_appends(self):
        existing = self.client.post('/api/meetings', json={'title': self.packet['meeting']['title'] + '专题', 'meeting_date': '2026-07-15', 'key_discussions': '原有讨论'}).json
        self.decision['meeting_id'] = existing['id']
        self.assertEqual(self.commit().status_code, 201)
        meeting = self.rows('meeting')[0]
        self.assertEqual(meeting['title'], existing['title'])
        self.assertEqual(meeting['key_discussions'], '原有讨论\n\n循环寿命与性能评价')

    def test_same_day_unrelated_and_similar_other_day_are_candidates(self):
        unrelated = self.client.post('/api/meetings', json={'title': '车辆照明专题', 'meeting_date': '2026-07-15'}).json
        similar = self.client.post('/api/meetings', json={'title': self.packet['meeting']['title'], 'meeting_date': '2026-07-16'}).json
        omitted = self.client.post('/api/meetings', json={'title': '材料学讲座', 'meeting_date': '2026-08-01'}).json
        candidates_by_id = {m['id']: m for m in self.preview().json['meetings']}
        self.assertTrue(candidates_by_id[unrelated['id']]['same_day'])
        self.assertFalse(candidates_by_id[unrelated['id']]['similar_title'])
        self.assertTrue(candidates_by_id[similar['id']]['similar_title'])
        self.assertFalse(candidates_by_id[similar['id']]['same_day'])
        self.assertNotIn(omitted['id'], candidates_by_id)
        self.decision['meeting_id'] = similar['id']
        self.assertEqual(self.commit().status_code, 201)
        self.assertEqual(self.client.get(f'/api/meetings/{similar["id"]}').json['meeting_date'], '2026-07-16')

    def reimport_setup(self):
        result = self.commit()
        self.assertEqual(result.status_code, 201, result.json)
        self.decision['meeting_id'] = result.json['meeting_id']
        self.decision['standards']['S1'] = {'mode': 'existing', 'id': self.rows('standard')[0]['id']}
        return self.rows('action_item')[0]

    def action_plan(self, review):
        return {'existing_snapshot': review['existing_snapshot'],
                'keep_existing': [a['id'] for a in review['existing_snapshot']], 'keep_new': []}

    def test_review_lists_new_old_and_defaults_duplicate_to_omit(self):
        old = self.reimport_setup()
        before = self.rows('action_item')
        self.packet['actions'][0]['description'] = '根据最新意见修订草案'
        r = self.review()
        self.assertEqual(r.status_code, 200, r.json)
        self.assertEqual(len(r.json['rows']), 2)
        existing, new = r.json['rows']
        self.assertEqual(existing['key'], f'existing:{old["id"]}')
        self.assertTrue(existing['keep'])
        self.assertFalse(new['keep'])
        self.assertIn(new['key'], existing['duplicates'])
        self.assertEqual(self.rows('action_item'), before)

    def test_user_can_replace_existing_action_with_new_and_remove_history(self):
        old = self.reimport_setup()
        self.packet['actions'][0]['description'] = '更新后的任务说明'
        plan = self.action_plan(self.review().json)
        plan.update(keep_existing=[], keep_new=[0])
        r = self.commit(plan)
        self.assertEqual(r.status_code, 201, r.json)
        self.assertEqual(r.json['deleted_actions'], 1)
        self.assertEqual(r.json['created_actions'], 1)
        self.assertEqual(len(self.rows('action_item')), 1)
        self.assertIn('更新后的任务说明', self.rows('action_item')[0]['description'])
        self.assertNotEqual(self.rows('action_item')[0]['id'], old['id'])
        self.assertEqual(len(self.rows('action_status_history')), 1)

    def test_user_can_explicitly_keep_both_duplicates(self):
        self.reimport_setup()
        plan = self.action_plan(self.review().json)
        plan['keep_new'] = [0]
        r = self.commit(plan)
        self.assertEqual(r.status_code, 201, r.json)
        self.assertEqual(len(self.rows('action_item')), 2)

    def test_deleted_new_action_never_creates_history_or_owner(self):
        self.packet['actions'][0]['stakeholder'] = 'Not Imported'
        plan = self.action_plan(self.review().json)
        r = self.commit(plan)
        self.assertEqual(r.status_code, 201, r.json)
        self.assertEqual(self.rows('action_item'), [])
        self.assertEqual(self.rows('action_status_history'), [])
        self.assertFalse(any(r['value']=='Not Imported' for r in self.rows('lookup_value')))

    def test_delete_old_collect_comments_cascades_feedback_and_history(self):
        self.reimport_setup()
        created = self.client.post('/api/actions', json={'item_type': 'Collect Comments',
            'standard_id': self.rows('standard')[0]['id'], 'meeting_id': self.decision['meeting_id'],
            'title': '收集试验意见', 'description': '向试验组征集意见',
            'recipients': [{'respondent_person': 'Tester'}]})
        self.assertEqual(created.status_code, 201, created.json)
        r = self.review().json
        row = next(a for a in r['rows'] if a['key']==f'existing:{created.json["id"]}')
        self.assertEqual(row['recipient_count'], 1)
        plan = self.action_plan(r)
        plan['keep_existing'].remove(created.json['id'])
        result = self.commit(plan)
        self.assertEqual(result.status_code, 201, result.json)
        self.assertEqual(self.rows('feedback_recipient'), [])
        self.assertEqual(self.rows('recipient_status_history'), [])
        self.assertEqual(len(self.rows('action_status_history')), 1)

    def test_stale_or_foreign_action_selection_is_rejected_atomically(self):
        old = self.reimport_setup()
        plan = self.action_plan(self.review().json)
        plan['keep_existing'] = [9999]
        self.assertEqual(self.commit(plan).status_code, 400)
        plan = self.action_plan(self.review().json)
        updated = self.client.put(f'/api/actions/{old["id"]}', json={'description': '其他窗口修改', 'coordinator': 'Tester'})
        self.assertEqual(updated.status_code, 200, updated.json)
        self.assertEqual(self.commit(plan).status_code, 409)
        self.assertEqual(len(self.rows('action_item')), 1)
        self.assertEqual(self.rows('action_item')[0]['description'], '其他窗口修改')

    def test_same_name_date_prevents_duplicate_meeting(self):
        self.assertEqual(self.commit().status_code, 201)
        self.assertEqual(self.commit().status_code, 400)
        self.assertEqual(len(self.rows('meeting')), 1)

    def test_reference_can_be_promoted_and_actions_relinked(self):
        result = self.commit().json
        self.decision['meeting_id'] = result['meeting_id']
        self.decision['standards']['S1'] = {'mode': 'existing', 'id': self.rows('standard')[0]['id']}
        self.decision['standards']['S2'] = {'mode': 'create', 'profile': {'name_cn': self.packet['standards'][1]['name_cn'], 'std_no': 'TEST-2'}}
        response = self.commit()
        self.assertEqual(response.status_code, 201, response.json)
        self.assertEqual(len(self.rows('meeting_standard')), 2)
        self.assertTrue(all(a['standard_id'] for a in self.rows('action_item')))

    def test_invalid_packets_are_400_and_write_nothing(self):
        original = copy.deepcopy(self.packet)
        mutations = [lambda p: p.update(actions={}), lambda p: p['meeting'].update(meeting_date='2026-02-30'),
                     lambda p: p['meeting'].update(title=2), lambda p: p['meeting'].update(meeting_date='20260715'),
                     lambda p: p['actions'][0].update(standard_key='missing'), lambda p: p['standards'].append(p['standards'][0]),
                     lambda p: p['standards'][0].update(extra='unknown'), lambda p: p['meeting'].update(key_discussions='')]
        for mutate in mutations:
            self.packet = copy.deepcopy(original)
            mutate(self.packet)
            response = self.client.post('/api/meetings/imports', json={'package': self.packet, 'decision': self.decision})
            self.assertEqual(response.status_code, 400, self.packet)
        self.assertEqual(self.rows('meeting'), [])

    def test_reference_material_fixture(self):
        self.packet = read_workbook(FileStorage(stream=io.BytesIO((ROOT / 'docs' / 'meeting-import-example.xlsx').read_bytes()), filename='example.xlsx'))
        self.decision['standards'] = {s['key']: {'mode': 'skip'} for s in self.packet['standards']}
        result = self.commit()
        self.assertEqual(result.status_code, 201, result.json)
        self.assertEqual(self.rows('meeting')[0]['meeting_date'], '2026-07-15')
        self.assertEqual(len(self.rows('meeting_standard')), 0)
        self.assertEqual(len(self.rows('action_item')), 0)
        self.assertEqual(result.json['excluded_actions'], 5)
        self.assertEqual(self.rows('standard'), [])

    def test_exchange_contains_only_tracked_minutes(self):
        self.assertEqual(self.commit().status_code, 201)
        with self.client.get('/api/transfer/export?user=Tester') as response:
            bundle = response.data
        tables = json.loads(bundle)['tables']
        self.assertNotIn('meeting_reference', tables)
        other_path = new_database()
        init_db(other_path, with_seed=False)
        other = create_app(other_path).test_client()
        inspected = other.post('/api/transfer/inspect', data={'file': (io.BytesIO(bundle), 'exchange.json')})
        self.assertEqual(inspected.status_code, 200, inspected.json)
        imported = other.post('/api/transfer/import', data={'file': (io.BytesIO(bundle), 'exchange.json'), 'as_user': 'Receiver'})
        self.assertEqual(imported.status_code, 200, imported.json)
        meeting = other.get('/api/meetings').json['items'][0]
        self.assertNotIn('references', meeting)
        self.assertEqual(len(meeting['standards']), 1)

    def test_archived_and_wrong_date_choices_are_rejected(self):
        first = self.commit().json
        sid = self.rows('standard')[0]['id']
        self.decision['meeting_id'] = first['meeting_id']
        self.decision['standards']['S1'] = {'mode': 'existing', 'id': sid}
        self.client.post(f'/api/standards/{sid}/archive')
        self.assertEqual(self.commit().status_code, 400)
        self.packet['meeting']['meeting_date'] = '2026-07-16'
        self.assertEqual(self.commit().status_code, 400)

    def test_confirmation_and_empty_optional_lists(self):
        self.packet['standards'] = []
        self.packet['actions'] = []
        self.decision['standards'] = {}
        self.decision['confirmed'] = False
        self.assertEqual(self.commit().status_code, 400)
        self.assertEqual(self.rows('meeting'), [])
        self.decision['confirmed'] = True
        self.assertEqual(self.commit().status_code, 201)
        self.assertEqual(self.rows('standard'), [])

    def test_all_untracked_writes_no_actions_or_history_or_owners(self):
        self.packet['actions'][0]['stakeholder'] = 'Excluded Owner'
        self.decision['standards'] = {s['key']: {'mode': 'skip'} for s in self.packet['standards']}
        response = self.commit()
        self.assertEqual(response.status_code, 201, response.json)
        self.assertEqual(response.json['excluded_actions'], 2)
        self.assertEqual(self.rows('action_item'), [])
        self.assertEqual(self.rows('action_status_history'), [])
        self.assertFalse(any(r['value'] == 'Excluded Owner' for r in self.rows('lookup_value')))
        self.assertEqual(self.client.get('/api/actions').json['total'], 0)

    def test_shared_and_meeting_wide_actions(self):
        self.packet['actions'][1]['standard_key'] = 'S1,S2'
        self.packet['actions'].append({'title': '安排下次会议', 'description': '预订会议室', 'standard_key': ''})
        response = self.commit()
        self.assertEqual(response.status_code, 201, response.json)
        self.assertEqual(response.json['created_actions'], 3)
        self.assertIsNotNone(self.rows('action_item')[1]['standard_id'])
        self.assertIsNone(self.rows('action_item')[2]['standard_id'])
        self.assertNotIn(self.packet['standards'][1]['name_cn'], self.rows('action_item')[1]['description'])

    def test_excel_validation_and_removed_json_upload(self):
        for content, filename in [(b'not excel', 'bad.xlsx'), (b'{}', 'bad.json'), (b'x' * (2*1024*1024+1), 'big.xlsx')]:
            r = self.client.post('/api/meetings/imports/preview', data={'file': (io.BytesIO(content), filename)})
            self.assertEqual(r.status_code, 400, r.json)
        r = self.client.post('/api/meetings/imports/preview', json={'package': self.packet})
        self.assertEqual(r.status_code, 400)
        for mutate in [lambda w: setattr(w['Meeting']['A1'], 'value', 'wrong header'),
                       lambda w: setattr(w['Actions']['A3'], 'value', '=1+1'),
                       lambda w: setattr(w['Actions']['C3'], 'value', 'S99'),
                       lambda w: w.create_sheet('Extra'),
                       lambda w: setattr(w['Meeting']['A4'], 'value', 'second meeting')]:
            workbook = load_workbook(io.BytesIO(self.excel()))
            mutate(workbook)
            stream = io.BytesIO(); workbook.save(stream); workbook.close()
            r = self.client.post('/api/meetings/imports/preview', data={'file': (io.BytesIO(stream.getvalue()), 'bad.xlsx')})
            self.assertEqual(r.status_code, 400, r.json)
        self.assertEqual(self.rows('meeting'), [])

    def test_excel_dates_and_reference_fixture_roundtrip(self):
        content = (ROOT / 'docs' / 'meeting-import-example.xlsx').read_bytes()
        r = self.client.post('/api/meetings/imports/preview', data={'file': (io.BytesIO(content), 'example.xlsx')})
        self.assertEqual(r.status_code, 200, r.json)
        self.assertEqual(r.json['package']['meeting']['meeting_date'], '2026-07-15')
        self.assertEqual(r.json['package']['actions'][1]['standard_key'], 'S2,S3')
        self.assertEqual(r.json['package']['actions'][0]['due_date'], '2026-09-25')


    def test_current_stage_is_required_and_exact_enum(self):
        for value in ['', '实施', '送审稿', 'DRAFTING', '不存在']:
            self.packet['standards'][0]['current_stage'] = value
            response = self.commit()
            self.assertEqual(response.status_code, 400, response.json)
            self.assertEqual(self.rows('meeting'), [])
        self.packet['standards'][0].pop('current_stage')
        self.assertEqual(self.commit().status_code, 400)

    def test_all_meeting_stages_can_create_profiles(self):
        for stage in ['预研', '立项', '起草', '征求意见', '审查', '报批', '发布', '复审', '废止']:
            self.packet['standards'][0]['current_stage'] = stage
            self.packet['standards'][0]['name_cn'] = '阶段测试标准' + stage
            self.packet['meeting']['title'] = '阶段测试会议' + stage
            self.decision['standards']['S1']['profile'] = {'name_cn': '阶段测试标准' + stage, 'std_no': ''}
            response = self.commit()
            self.assertEqual(response.status_code, 201, response.json)
        self.assertEqual(len(self.rows('standard_stage_history')), 9)
        self.assertEqual({row['effective_date'] for row in self.rows('standard_stage_history')}, {'2026-07-15'})

    def test_reimport_replaces_meeting_stage_snapshot(self):
        first = self.commit().json
        sid = self.rows('standard')[0]['id']
        self.decision['meeting_id'] = first['meeting_id']
        self.decision['standards']['S1'] = {'mode': 'existing', 'id': sid}
        self.packet['standards'][0]['current_stage'] = '审查'
        response = self.commit()
        self.assertEqual(response.status_code, 201, response.json)
        self.assertEqual(response.json['advanced_standards'], 1)
        self.assertEqual(self.rows('standard')[0]['stage_code'], 'REVIEW_DRAFT')
        self.assertEqual(len(self.rows('standard_stage_history')), 2)
        record = self.rows('standard_stage_history')[-1]
        self.assertEqual(record['created_by'], 'Tester')
        self.assertEqual(record['effective_date'], '2026-07-15')
        self.assertIn(self.packet['meeting']['title'], record['reference'])
        self.assertEqual(self.commit().json['advanced_standards'], 0)
        self.packet['standards'][0]['current_stage'] = '起草'
        self.assertEqual(self.commit().status_code, 201)
        self.assertEqual(self.rows('standard')[0]['stage_code'], 'DRAFTING')
        self.assertEqual(len(self.rows('standard_stage_history')), 2)

    def test_stage_advancement_rolls_back_with_later_invalid_choice(self):
        first = self.commit().json
        sid = self.rows('standard')[0]['id']
        self.decision['meeting_id'] = first['meeting_id']
        self.decision['standards']['S1'] = {'mode': 'existing', 'id': sid}
        self.packet['standards'][0]['current_stage'] = '发布'
        self.packet['meeting']['key_discussions'] = '不应保存的新议题'
        self.decision['standards']['S2'] = {'mode': 'invalid'}
        response = self.commit()
        self.assertEqual(response.status_code, 400, response.json)
        self.assertEqual(self.rows('standard')[0]['stage_code'], 'DRAFTING')
        self.assertEqual(len(self.rows('standard_stage_history')), 1)
        self.assertNotIn('不应保存', self.rows('meeting')[0]['key_discussions'])

    def test_older_higher_stage_meeting_imports_without_overriding_latest(self):
        first = self.commit().json
        sid = self.rows('standard')[0]['id']
        self.decision['standards']['S1'] = {'mode': 'existing', 'id': sid}
        self.packet['meeting']['title'] = '更早会议'
        self.packet['meeting']['meeting_date'] = '2026-01-01'
        self.packet['standards'][0]['current_stage'] = '报批'
        response = self.commit()
        self.assertEqual(response.status_code, 201, response.json)
        self.assertEqual(self.rows('standard')[0]['stage_code'], 'DRAFTING')
        self.assertEqual(len(self.rows('meeting')), 2)
        self.assertEqual(len(self.rows('standard_stage_history')), 2)

    def test_all_event_orders_use_latest_meeting_and_manual_max(self):
        for index, order in enumerate(permutations(['old', 'latest', 'manual'])):
            with self.subTest(order=order):
                name = f'乱序标准{index}'
                created = self.client.post('/api/standards', json={
                    'name_cn': name, 'stage_code': 'PRE_RESEARCH',
                    'stage_effective_date': '2026-01-01'})
                self.assertEqual(created.status_code, 201, created.json)
                sid = created.json['id']
                self.packet['standards'][0]['name_cn'] = name
                self.decision['standards']['S1'] = {'mode': 'existing', 'id': sid}
                self.decision['meeting_id'] = 'new'
                seen = set()
                for event in order:
                    if event == 'manual':
                        response = self.client.post(f'/api/standards/{sid}/stages', json={
                            'stage_code': 'REVIEW_DRAFT', 'effective_date': '2026-12-01'})
                        self.assertEqual(response.status_code, 201, response.json)
                    else:
                        self.packet['meeting']['title'] = f'乱序会议{index}{event}'
                        self.packet['meeting']['meeting_date'] = '2026-03-01' if event == 'old' else '2026-07-01'
                        self.packet['standards'][0]['current_stage'] = '发布' if event == 'old' else '起草'
                        response = self.commit()
                        self.assertEqual(response.status_code, 201, response.json)
                    seen.add(event)
                    expected = 'DRAFTING' if 'latest' in seen else 'RELEASED' if 'old' in seen else 'PRE_RESEARCH'
                    if 'manual' in seen and expected in ('DRAFTING', 'PRE_RESEARCH'):
                        expected = 'REVIEW_DRAFT'
                    detail = self.client.get(f'/api/standards/{sid}').json
                    self.assertEqual(detail['stage_code'], expected)
                    lifecycle = self.client.get(f'/api/standards/{sid}/lifecycle').json
                    current = [row['code'] for row in lifecycle['stages'] if row['state'] == 'current']
                    self.assertEqual(current, [expected])

    def test_newer_meeting_can_exceed_manual_stage(self):
        first = self.commit().json
        sid = self.rows('standard')[0]['id']
        manual = self.client.post(f'/api/standards/{sid}/stages', json={
            'stage_code': 'REVIEW_DRAFT', 'effective_date': '2026-07-16'})
        self.assertEqual(manual.status_code, 201, manual.json)
        self.decision['standards']['S1'] = {'mode': 'existing', 'id': sid}
        self.packet['meeting']['title'] = '发布会议'
        self.packet['meeting']['meeting_date'] = '2026-08-01'
        self.packet['standards'][0]['current_stage'] = '发布'
        self.assertEqual(self.commit().status_code, 201)
        self.assertEqual(self.client.get(f'/api/standards/{sid}').json['stage_code'], 'RELEASED')

    def test_date_edit_detach_and_delete_recompute_sources(self):
        first = self.commit().json
        sid = self.rows('standard')[0]['id']
        self.decision['standards']['S1'] = {'mode': 'existing', 'id': sid}
        self.packet['meeting']['title'] = '更早发布会议'
        self.packet['meeting']['meeting_date'] = '2026-03-01'
        self.packet['standards'][0]['current_stage'] = '发布'
        older = self.commit().json
        response = self.client.put(f'/api/meetings/{older["meeting_id"]}', json={'meeting_date': '2026-09-01'})
        self.assertEqual(response.status_code, 200, response.json)
        self.assertEqual(self.client.get(f'/api/standards/{sid}').json['stage_code'], 'RELEASED')
        response = self.client.delete(f'/api/meetings/{older["meeting_id"]}/standards/{sid}?drop_notes=1')
        self.assertEqual(response.status_code, 200, response.json)
        self.assertEqual(self.client.get(f'/api/standards/{sid}').json['stage_code'], 'DRAFTING')
        self.assertEqual(self.client.delete(f'/api/meetings/{first["meeting_id"]}').status_code, 200)
        self.assertEqual(self.client.get(f'/api/standards/{sid}').json['stage_code'], 'PRE_RESEARCH')
        detail = self.client.get(f'/api/standards/{sid}').json
        self.assertEqual([s['code'] for s in detail['stage_timeline'] if s['state'] == 'current'],
                         ['PRE_RESEARCH'])

    def test_same_day_meetings_use_max_stage_independent_of_import_order(self):
        first = self.commit().json
        sid = self.rows('standard')[0]['id']
        self.decision['standards']['S1'] = {'mode': 'existing', 'id': sid}
        self.packet['meeting']['title'] = '同日发布会议'
        self.packet['standards'][0]['current_stage'] = '发布'
        self.assertEqual(self.commit().status_code, 201)
        self.assertEqual(self.client.get(f'/api/standards/{sid}').json['stage_code'], 'RELEASED')
        self.decision['meeting_id'] = first['meeting_id']
        self.packet['meeting']['title'] = '动力蓄电池性能工作组第二次会议'
        self.packet['standards'][0]['current_stage'] = '立项'
        self.assertEqual(self.commit().status_code, 201)
        self.assertEqual(self.client.get(f'/api/standards/{sid}').json['stage_code'], 'RELEASED')

    def test_edit_imported_stage_becomes_manual_floor(self):
        self.packet['standards'][0]['current_stage'] = '发布'
        first = self.commit().json
        sid = self.rows('standard')[0]['id']
        record = self.rows('standard_stage_history')[0]
        self.assertEqual(record['is_manual'], 0)
        response = self.client.put(f'/api/standards/{sid}/stages/{record["id"]}', json={
            'effective_date': '2026-07-15', 'note': '人工确认发布阶段'})
        self.assertEqual(response.status_code, 200, response.json)
        self.assertEqual(self.rows('standard_stage_history')[0]['is_manual'], 1)
        self.decision['standards']['S1'] = {'mode': 'existing', 'id': sid}
        self.packet['meeting']['title'] = '后续起草会议'
        self.packet['meeting']['meeting_date'] = '2026-08-01'
        self.packet['standards'][0]['current_stage'] = '起草'
        self.assertEqual(self.commit().status_code, 201)
        self.assertEqual(self.client.get(f'/api/standards/{sid}').json['stage_code'], 'RELEASED')

    def test_exchange_preserves_stage_sources_and_recomputes_current(self):
        self.packet['standards'][0]['current_stage'] = '发布'
        self.packet['meeting']['meeting_date'] = '2026-03-01'
        self.assertEqual(self.commit().status_code, 201)
        sid = self.rows('standard')[0]['id']
        self.decision['standards']['S1'] = {'mode': 'existing', 'id': sid}
        self.packet['meeting']['title'] = '最新起草会议'
        self.packet['meeting']['meeting_date'] = '2026-08-01'
        self.packet['standards'][0]['current_stage'] = '起草'
        self.assertEqual(self.commit().status_code, 201)
        exported = self.client.get('/api/transfer/export?user=Tester')
        self.assertEqual(exported.status_code, 200)
        package = json.loads(exported.data)
        # Cached stage is deliberately wrong: import must derive it from source records.
        package['tables']['standard'][0]['stage_code'] = 'ABOLISHED'
        other_path = new_database(initialize=True)
        other = create_app(other_path).test_client()
        response = other.post('/api/transfer/import', data={
            'file': (io.BytesIO(json.dumps(package).encode()), 'data.json'), 'as_user': 'Tester'})
        self.assertEqual(response.status_code, 200, response.json)
        detail = other.get('/api/standards/1').json
        self.assertEqual(detail['stage_code'], 'DRAFTING')
        self.assertTrue(all(row['is_manual'] == 0 for row in detail['stage_history']))

    def test_implemented_stage_is_manual_and_can_advance_to_reexamination(self):
        self.packet['standards'][0]['current_stage'] = '发布'
        first = self.commit().json
        sid = self.rows('standard')[0]['id']
        response = self.client.post(f'/api/standards/{sid}/stages', json={
            'stage_code': 'IMPLEMENTED', 'effective_date': '2026-07-16'}, headers={'X-User': 'Tester'})
        self.assertEqual(response.status_code, 201, response.json)
        self.decision['standards']['S1'] = {'mode': 'existing', 'id': sid}
        self.packet['meeting']['title'] = '复审会议'
        self.packet['meeting']['meeting_date'] = '2026-08-01'
        self.packet['standards'][0]['current_stage'] = '复审'
        response = self.commit()
        self.assertEqual(response.status_code, 201, response.json)
        self.assertEqual(self.rows('standard')[0]['stage_code'], 'REEXAMINATION')
        codes = [s['code'] for s in self.client.get('/api/meta').json['stages']]
        self.assertEqual(codes[codes.index('RELEASED') + 1:codes.index('ABOLISHED')], ['IMPLEMENTED', 'REEXAMINATION'])


if __name__ == '__main__':
    unittest.main()
