"""意见交换 Excel：固定五列，导出、模板与导入读取共用一份定义。

    序号 | 章条编号 | 章节内容 | 意见内容（包括理由或依据） | 意见状态

只有「章条编号」「意见内容」「意见状态」参与导入；「序号」「章节内容」仅供阅读。
"""

import math
import unicodedata
from io import BytesIO
from zipfile import BadZipFile
from xml.etree.ElementTree import ParseError

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl.worksheet.datavalidation import DataValidation

from common import ApiError
from constants import COMMENT_STATUSES

HEADERS = ("序号", "章条编号", "章节内容", "意见内容（包括理由或依据）", "意见状态")
SHEET = "意见"
MAX_ROWS = 1000
# 列宽上下限（字符单位）：长文本列换行显示，其余列按内容收紧。
MAX_WIDTH = (8, 14, 60, 60, 20)
MIN_WIDTH = (6, 10, 24, 30, 14)
LINE_HEIGHT = 15

_HEAD_FILL = PatternFill("solid", fgColor="1F5C8B")
_HEAD_FONT = Font(bold=True, color="FFFFFF")
_RULE = Side(style="thin", color="C9D3DE")
_BORDER = Border(left=_RULE, right=_RULE, top=_RULE, bottom=_RULE)


def _display_width(text: str) -> int:
    """全角字符按两个字符宽度计算，与 Excel 列宽单位大致对应。"""
    return sum(2 if unicodedata.east_asian_width(ch) in "WF" else 1 for ch in text)


def _line_count(text: str, width: float) -> int:
    usable = max(1.0, width - 1)
    return sum(max(1, math.ceil(_display_width(line) / usable)) for line in text.split("\n"))


def build(rows: list[tuple]) -> BytesIO:
    """生成工作簿；rows 为空时即导入模板。每行为 (章条编号, 章节内容, 意见内容, 意见状态)。"""
    book = Workbook()
    sheet = book.active
    sheet.title = SHEET
    sheet.append(HEADERS)
    for number, row in enumerate(rows, start=1):
        sheet.append((number, *row))

    widths = []
    for index, header in enumerate(HEADERS):
        longest = max([_display_width(header)] + [
            max(_display_width(line) for line in str(row[index - 1] or "").split("\n"))
            for row in rows if index])
        widths.append(min(MAX_WIDTH[index], max(MIN_WIDTH[index], longest + 2)))
        sheet.column_dimensions[get_column_letter(index + 1)].width = widths[-1]

    for cell in sheet[1]:
        cell.fill, cell.font, cell.border = _HEAD_FILL, _HEAD_FONT, _BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    sheet.row_dimensions[1].height = LINE_HEIGHT * max(
        _line_count(header, width) for header, width in zip(HEADERS, widths)) + 6

    for row in sheet.iter_rows(min_row=2):
        lines = 1
        for index, cell in enumerate(row):
            cell.border = _BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=True,
                                       horizontal="center" if index in (0, 1, 4) else "left")
            if cell.value is not None:
                lines = max(lines, _line_count(str(cell.value), widths[index]))
        sheet.row_dimensions[row[0].row].height = LINE_HEIGHT * lines + 4

    # 章条编号按文本保存，避免 Excel 把 5.10 改写成 5.1。
    sheet.column_dimensions["B"].number_format = "@"
    for (cell,) in sheet.iter_rows(min_row=2, min_col=2, max_col=2):
        cell.number_format = "@"
    status = DataValidation(type="list", formula1='"' + ",".join(COMMENT_STATUSES) + '"',
                            allow_blank=False, showErrorMessage=True,
                            errorTitle="意见状态", error="请从下拉列表中选择意见状态")
    status.add(f"E2:E{MAX_ROWS + 1}")
    sheet.add_data_validation(status)
    sheet.freeze_panes = "A2"

    guide = book.create_sheet("使用说明")
    guide.append(("列名", "是否导入", "填写方法", "示例"))
    guide_rows = [
        (HEADERS[0], "否", "仅供阅读和排序，不是系统意见编号，可留空。", "1"),
        (HEADERS[1], "是", "填写所选草案中存在的章条编号；留空表示整份草案。建议使用文本格式，避免 Excel 自动转为日期。", "5.1"),
        (HEADERS[2], "否", "原文参考，不会覆盖草案正文。", "对应章条原文"),
        (HEADERS[3], "是，必填", "每行一条意见，写明建议修改内容及理由或依据。", "建议修改本条表述，理由：…"),
        (HEADERS[4], "是，必填", "从下拉列表选择，允许值：" + " / ".join(COMMENT_STATUSES), "Draft"),
        ("导入规则", "", "保留第一张‘意见’工作表的五列表头和顺序。导入时在页面选择目标草案及共用元数据。任一行校验失败时整批不写入。", ""),
    ]
    for row in guide_rows:
        guide.append(row)
    for cell in guide[1]:
        cell.fill = _HEAD_FILL
        cell.font = _HEAD_FONT
    for row in guide.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        guide.row_dimensions[row[0].row].height = 64
    for col, width in zip("ABCD", (30, 18, 88, 36)):
        guide.column_dimensions[col].width = width
    guide.freeze_panes = "A2"

    out = BytesIO()
    book.save(out)
    out.seek(0)
    return out


def _text(cell) -> str:
    value = cell.value
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def read(upload, t) -> list[dict]:
    """读取并逐行校验上传文件；任何一行有误都整体拒绝，不写库。t(zh, en) 选择提示语言。"""
    if not upload or not upload.filename.lower().endswith(".xlsx"):
        raise ApiError(t("请上传 .xlsx 意见文件，可先下载导入模板", "Upload an .xlsx opinion file; download the import template first"))
    try:
        book = load_workbook(upload.stream, read_only=False, data_only=False, keep_links=False)
    except (BadZipFile, InvalidFileException, ValueError, KeyError, OSError, ParseError):
        raise ApiError(t("无法读取此 Excel，请使用导入模板另存为 .xlsx 后上传",
                         "This Excel file cannot be read. Save it from the import template as .xlsx and try again")) from None
    try:
        sheet = book.worksheets[0]
        headers = tuple(_text(cell) for cell in sheet[1][:len(HEADERS)])
        if headers != HEADERS:
            raise ApiError(t("第 1 行表头与模板不一致，请保留五个固定字段及其顺序",
                             "Row 1 does not match the template. Keep the five fixed headers in order"),
                           detail=" | ".join(HEADERS))
        if sheet.max_row > MAX_ROWS + 1:
            raise ApiError(t(f"单次最多导入 {MAX_ROWS} 条意见", f"Import up to {MAX_ROWS} opinions at a time"))
        rows, errors = [], []
        for row in sheet.iter_rows(min_row=2, max_col=len(HEADERS)):
            number = row[0].row
            if any(cell.data_type in ("f", "e") for cell in row[1:]):
                errors.append(t(f"第 {number} 行含公式或错误值，请改为实际文本",
                                f"Row {number} contains a formula or error value; enter plain text"))
                continue
            clause, text, status = _text(row[1]), _text(row[3]), _text(row[4])
            if not (clause or text or status):
                continue
            if not text:
                errors.append(t(f"第 {number} 行缺少意见内容", f"Row {number}: opinion text is missing"))
            if status not in COMMENT_STATUSES:
                errors.append(t(f"第 {number} 行意见状态「{status or '空'}」不在可选值内",
                                f"Row {number}: status \"{status or 'blank'}\" is not an allowed value"))
            rows.append({"row": number, "clause_no": clause, "text": text, "status": status})
        if errors:
            raise ApiError(t(f"导入文件有 {len(errors)} 处问题，未写入任何数据",
                             f"The file has {len(errors)} problem(s). Nothing was imported"),
                           detail="\n".join(errors))
        if not rows:
            raise ApiError(t("文件中没有可导入的意见", "The file contains no opinions to import"))
        return rows
    finally:
        book.close()
