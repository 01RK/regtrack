"""Current business and interface regression coverage."""
from pathlib import Path
from openpyxl import load_workbook
from werkzeug.datastructures import FileStorage
import comment_workbook

ROOT = Path(__file__).resolve().parents[1]


def test_template_has_column_instructions():
    book = load_workbook(comment_workbook.build([]))
    assert book.sheetnames == [comment_workbook.SHEET, "\u4f7f\u7528\u8bf4\u660e"]
    assert tuple(c.value for c in book.worksheets[0][1]) == comment_workbook.HEADERS
    assert book.worksheets[1].max_row == 7
    for header in comment_workbook.HEADERS:
        assert header in [row[0] for row in book.worksheets[1].values]
    book.close()


def test_instruction_sheet_does_not_change_import():
    upload = FileStorage(stream=comment_workbook.build([("5.1", "source", "proposal", "Submitted")]), filename="opinions.xlsx")
    rows = comment_workbook.read(upload, lambda zh, en: en)
    assert len(rows) == 1
    assert rows[0]["text"] == "proposal"
    assert rows[0]["status"] == "Submitted"


def test_readonly_metadata_has_no_edit_trigger():
    source = (ROOT / "static/js/inline-edit.js").read_text()
    assert "cell.onclick" not in source
    assert "cell.onkeydown" not in source
    assert 'role="button"' not in source


def test_action_status_is_title_badge():
    template = (ROOT / "templates/actions.html").read_text()
    assert template.count("<col ") == 6
    assert "action-col-status" not in template
    source = (ROOT / "static/js/pages/actions.js").read_text()
    assert "action-title-status" in source
    assert "emptyRow(6," in source


def test_lifecycle_review_has_explicit_status_filter_and_equal_columns():
    source = (ROOT / "static/js/pages/lifecycle.js").read_text()
    assert 'new Set(["Accepted", "Partially Accepted", "Rejected", "Submitted"])' in source
    assert "REVIEW_STATUSES.has(c.status)" in source
    assert "repeat(${stages.length}, minmax(280px, 1fr))" in source
    assert "${esc(m.meeting_no)}" not in source
    assert "${esc(c.comment_no)}" not in source


def test_comparison_uses_soft_line_breaks():
    source = (ROOT / "static/js/pages/comparison-detail.js").read_text()
    assert "RichText.render(markdown, {breaks: false})" in source
    base = (ROOT / "templates/base.html").read_text()
    assert "RegTrack {{ app_version }}" not in base
