"""Single comparison workbook: both source drafts, figures and the result sheet. Uses existing openpyxl."""
import io
import math
import re
import unicodedata
import zipfile
from xml.etree.ElementTree import ParseError
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl import Workbook, load_workbook
from openpyxl.drawing.image import Image
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
import db
import files
from common import ApiError
from comparison_service import HEADERS, META_HEADERS, CHANGE_TYPES, CLASSES, meta, chapters

SHEETS = ('Comparison', 'Changes', 'Before', 'After', 'Figures', 'Instructions')
SOURCE_HEADERS = ('Clause_no', 'Parent_clause_no', 'Level', 'Source_page', 'Title_cn', 'Content_cn',
                  'Title_en', 'Content_en', 'Illustration', 'Figures')
FIGURE_HEADERS = ('Figure_id', 'Side', 'Clause_no', 'Position', 'Image')
FIGURE_BOX = (420, 300)  # max width / height in pixels


def filename(b, a, kind):
    standard = db.query_one('SELECT std_no, name_cn FROM standard WHERE id=%s', (b['standard_id'],))
    name = (f"{standard['std_no'] or standard['name_cn']}_{b['version_name']}v{b['sub_version_no']}"
            f"_{a['version_name']}v{a['sub_version_no']}_{kind}.xlsx")
    return re.sub(r'[\\/:*?"<>|]+', '-', name)


def add_sources(wb, b, a):
    figures = wb['Figures']
    for side, draft in (('Before', b), ('After', a)):
        ws = wb[side]
        for c in chapters(draft['import_id']):
            ids = []
            for image in c['images']:
                fid = f'F{figures.max_row:03d}'
                ids.append(fid)
                figures.append([fid, side, c['clause_no'], image['position'], ''])
                picture = Image(io.BytesIO(files.read(image['file_path'])))
                scale = min(1, FIGURE_BOX[0] / picture.width, FIGURE_BOX[1] / picture.height)
                picture.width, picture.height = round(picture.width * scale), round(picture.height * scale)
                figures.add_image(picture, f'E{figures.max_row}')
                figures.row_dimensions[figures.max_row].height = picture.height * 0.75 + 8
            ws.append([c['clause_no'], c['parent_clause_no'] or '', c['level'], c['source_page'] or '',
                       c['title_cn'] or '', c['content_cn'] or '', c['title_en'] or '', c['content_en'] or '',
                       c['illustration'] or '', '; '.join(ids)])


def style(wb):
    for sheet in wb:
        sheet.freeze_panes = 'A2'
        sheet.sheet_view.showGridLines = False
        for cell in sheet[1]:
            cell.font = Font(color='FFFFFF', bold=True)
            cell.fill = PatternFill('solid', fgColor='243E58')
        sheet.row_dimensions[1].height = 28
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                if isinstance(cell.value, str):
                    cell.data_type = 's'
                cell.alignment = Alignment(vertical='top', wrap_text=True)
        for col in sheet.columns:
            sheet.column_dimensions[col[0].column_letter].width = 16
    for col in ('F', 'G', 'H', 'I', 'J'):
        wb['Changes'].column_dimensions[col].width = 54
    for side in ('Before', 'After'):
        for col, width in (('E', 24), ('F', 60), ('G', 24), ('H', 60), ('I', 30)):
            wb[side].column_dimensions[col].width = width
    wb['Figures'].column_dimensions['E'].width = FIGURE_BOX[0] / 7 + 2
    wb['Instructions'].column_dimensions['B'].width = 100
    for sheet in wb:
        if sheet.title == 'Figures':
            continue
        for row in sheet.iter_rows(min_row=2):
            lines = 1
            for cell in row:
                width = sheet.column_dimensions[cell.column_letter].width - 2
                text_lines = str(cell.value or '').split('\n')
                count = sum(max(1, math.ceil(sum(2 if unicodedata.east_asian_width(ch) in 'WF' else 1 for ch in line) / width)) for line in text_lines)
                lines = max(lines, count)
            sheet.row_dimensions[row[0].row].height = min(409, max(28, lines * 16 + 8))


def workbook(b, a, changes=()):
    wb = Workbook()
    ws = wb.active
    ws.title = 'Comparison'
    ws.append(META_HEADERS)
    ws.append(meta(b, a))
    changes_ws = wb.create_sheet('Changes')
    changes_ws.append(HEADERS)
    for c in changes:
        changes_ws.append([c['row_key'], ';'.join(c['before_nos']), ';'.join(c['after_nos']),
                           c['change_type'], c['revision_class'], c['summary_cn'], c['summary_en'],
                           c['reason_cn'], c['reason_en'], c['review_note']])
    for side in ('Before', 'After'):
        wb.create_sheet(side).append(SOURCE_HEADERS)
    wb.create_sheet('Figures').append(FIGURE_HEADERS)
    guide = wb.create_sheet('Instructions')
    guide.append(['Field', 'Rule / 填写规则'])
    for h, rule in zip(HEADERS, [
        'Unique stable row key, e.g. R001 / 唯一行标识',
        'Exact Clause_no from the Before sheet; separate with ; / 前版原始章节号，以英文分号分隔',
        'Exact Clause_no from the After sheet; separate with ; / 后版原始章节号，以英文分号分隔',
        ', '.join(CHANGE_TYPES), ', '.join(CLASSES),
        '具体变化的中文说明，不重复原文', 'Faithful English translation of Summary_cn',
        '分类依据，指出要求、条件、数值、适用范围或图示的变化',
        'Faithful English translation of Reason_cn',
        '可选：说明缺失证据或人工核实事项；建议中英双语',
    ]):
        guide.append([h, rule])
    guide.append(['Sources', 'Before、After 为两版原文，Figures 为附图（按 Figure_id 对应原文 Figures 列）。只读，不得修改。'])
    guide.append(['Coverage', '两侧每条原始章节均且仅覆盖一次，包括未变化条款；不得复制或改写原文到 Changes。'])
    guide.append(['Metadata', 'Comparison IDs and batch numbers are informational only; imports use the selected drafts. / 编号和批次仅供参考，导入以页面选择的草案为准。'])
    guide.append(['Analysis import', 'Analysis labels and notes are imported as supplied. No source-text equality or rationale completeness checks. / 分析分类和说明按文件导入，不强制检查原文一致性或依据完整性。'])
    add_sources(wb, b, a)
    style(wb)
    for col, values in [('D', CHANGE_TYPES), ('E', CLASSES)]:
        dv = DataValidation(type='list', formula1='"' + ','.join(values) + '"')
        changes_ws.add_data_validation(dv)
        dv.add(f'{col}2:{col}10001')
    wb.active = 1
    result = io.BytesIO()
    wb.save(result)
    result.seek(0)
    return result


def read_workbook(upload):
    if not upload or not upload.filename.lower().endswith('.xlsx'):
        raise ApiError('Upload an .xlsx workbook / 请上传 .xlsx 文件')
    data = upload.read()
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            if sum(i.file_size for i in archive.infolist()) > 80 * 1024 * 1024:
                raise ApiError('Expanded workbook is too large / 解压后的工作簿过大')
        wb = load_workbook(io.BytesIO(data), data_only=False, read_only=True)
    except (zipfile.BadZipFile, KeyError, ValueError, OSError, ParseError, InvalidFileException) as exc:
        raise ApiError('Invalid Excel workbook / 无效的 Excel 文件') from exc
    try:
        if set(wb.sheetnames) != set(SHEETS):
            raise ApiError('Keep all workbook sheets: ' + ', '.join(SHEETS) + ' / 必须保留比对工作簿的全部工作表')
        # Workbook source IDs are informational. The selected drafts own the import.
        changes = wb['Changes']
        if changes.max_row > 10001 or changes.max_column != len(HEADERS):
            raise ApiError('Changes supports up to 10000 rows and exactly 10 columns / 结果表最多10000行且必须为10列')
        if tuple(c.value for c in changes[1]) != HEADERS:
            raise ApiError('Changes headers do not match / 结果表表头不匹配')
        rows = []
        for row in changes.iter_rows(min_row=2):
            if any(c.data_type == 'f' for c in row):
                raise ApiError('Formulas are not allowed / 不允许使用公式')
            rows.append([c.value for c in row])
        return rows
    finally:
        wb.close()
