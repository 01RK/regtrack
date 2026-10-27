from testdb import new_database
"""意见导入 / 导出与界面语言切换。

    cd regtrack/codes && python -m unittest test_comment_exchange -v
"""

import io
import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent))

from app import create_app
from testdb import initialize as init_db  # noqa: E402
from dbtest import Database  # noqa: E402
from comment_workbook import HEADERS  # noqa: E402

ROOT = Path(__file__).resolve().parent.parent
HEADERS_USER = {"X-User": "Importer"}


def workbook(rows, headers=HEADERS):
    book = Workbook()
    sheet = book.active
    sheet.append(list(headers))
    for row in rows:
        sheet.append(row)
    out = io.BytesIO()
    book.save(out)
    out.seek(0)
    return out


class CommentExchangeTest(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.db = new_database()
        init_db(self.db, with_seed=True)
        self.client = create_app(self.db).test_client()

    def tearDown(self):
        self.tmp.cleanup()

    def upload(self, rows, draft_id=15, headers=HEADERS):
        return self.client.post("/api/comments/exchange/import", headers=HEADERS_USER,
                                data={"draft_id": str(draft_id), "file": (workbook(rows, headers), "x.xlsx")},
                                content_type="multipart/form-data")

    def test_template_has_fixed_headers_and_status_list(self):
        r = self.client.get("/api/comments/exchange/template")
        self.assertEqual(r.status_code, 200)
        sheet = load_workbook(io.BytesIO(r.data)).active
        self.assertEqual(tuple(c.value for c in sheet[1]), HEADERS)
        self.assertIn("Partially Accepted", sheet.data_validations.dataValidation[0].formula1)

    def test_export_numbers_rows_and_fills_chapter_text(self):
        r = self.client.get("/api/comments/exchange/export?draft_id=15")
        self.assertEqual(r.status_code, 200)
        self.assertIn("%E6%84%8F%E8%A7%81%E5%AF%BC%E5%87%BA", r.headers["Content-Disposition"])
        rows = [row for row in load_workbook(io.BytesIO(r.data)).active.iter_rows(min_row=2, values_only=True) if any(row)]
        self.assertEqual([row[0] for row in rows], list(range(1, len(rows) + 1)))
        clause_rows = [row for row in rows if row[1]]
        self.assertTrue(all(row[2] for row in clause_rows))

    def test_import_rejects_whole_file_on_unknown_clause_or_status(self):
        before = self.client.get("/api/comments?draft_id=15&page_size=50").json["total"]
        r = self.upload([[1, "1", "", "ok", "Accepted"], [2, "9.9", "", "bad", "Accepted"]])
        self.assertEqual(r.status_code, 400)
        self.assertIn("9.9", r.json["detail"])
        r = self.upload([[1, "1", "", "ok", "Unknown"]])
        self.assertEqual(r.status_code, 400)
        r = self.upload([[1, "1", "", "ok", "Accepted"]], headers=("A", "B", "C", "D", "E"))
        self.assertEqual(r.status_code, 400)
        self.assertEqual(self.client.get("/api/comments?draft_id=15&page_size=50").json["total"], before)

    def test_import_creates_single_draft_opinions_and_skips_duplicates(self):
        r = self.upload([[1, "1", "ignored", "Chapter opinion", "Rejected"],
                         [2, None, "", "Whole draft opinion", "Draft"]])
        self.assertEqual(r.status_code, 201, r.json)
        self.assertEqual(len(r.json["created"]), 2)
        with Database(self.db) as conn:
            rows = [tuple(row) for row in conn.execute("""SELECT c.clause_no, c.status, c.needs_follow_up, c.submitted_by, c.source_kind
                FROM comment c INNER JOIN draft_annotation a ON a.id = c.source_annotation_id
                WHERE a.created_by = 'Importer' ORDER BY c.id""").fetchall()]
        self.assertEqual(rows, [("1", "Rejected", 1, "Importer", "single"), (None, "Draft", 0, "Importer", "single")])
        again = self.upload([[1, "1", "", "Chapter opinion", "Accepted"]])
        self.assertEqual(again.json, {"created": [], "skipped_rows": [2]})

    def test_language_cookie_switches_page_labels(self):
        zh = self.client.get("/comments").get_data(as_text=True)
        self.assertIn('lang="zh-CN"', zh)
        self.assertIn('Comment Type<small class="cn-sub">意见类型</small>', zh)
        self.assertIn("导入意见", zh)
        self.client.set_cookie("regtrack_lang", "en")
        en = self.client.get("/comments").get_data(as_text=True)
        self.assertIn('lang="en"', en)
        self.assertIn("Import Opinions", en)
        # 英文界面表头只显示英文，不再输出中文小字
        self.assertIn('Comment Type', en)
        self.assertNotIn('意见类型', en)
        for page in ("/", "/standards", "/drafts", "/meetings", "/comments", "/actions", "/settings",
                     "/drafts/15", "/comments/import", "/comments/export", "/meetings/import", "/drafts/import"):
            self.assertNotIn('class="cn-sub"', self.client.get(page).get_data(as_text=True), page)
        self.assertIn("Import Draft Text", self.client.get("/drafts/15").get_data(as_text=True))


if __name__ == "__main__":
    unittest.main()
