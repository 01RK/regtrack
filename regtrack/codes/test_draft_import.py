from testdb import new_database
"""End-to-end coverage for chapter import, annotation editing and replacement."""
import io
import json
import tempfile
import unittest
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.drawing.image import Image
from PIL import Image as PillowImage

import db
from app import create_app
from testdb import initialize as init_db
from dbtest import STAMP, Database

TEMPLATE = Path(__file__).resolve().parents[1] / "static" / "standard-clause-import-blank.xlsx"


class DraftImportTest(unittest.TestCase):
    def setUp(self):
        self.folder = tempfile.TemporaryDirectory()
        self.db = new_database()
        init_db(self.db, with_seed=True)
        with Database(self.db) as conn:
            for sid, std_no, name in ((1, 'TEST 1-2026', '测试目标标准'),
                                      (32, 'GB 38031-2025', '电动汽车用动力蓄电池安全要求')):
                conn.insert("standard", dict(STAMP, id=sid, std_no=std_no, name_cn=name, stage_code='COMMENT_DRAFT',
                                             mb_owner='Reviewer', risk_level='High',
                                             created_by='Reviewer', updated_by='Reviewer'))
        self.client = create_app(self.db).test_client()
        book = load_workbook(TEMPLATE)
        sheet = book["Import_Data"]
        headers = [cell.value for cell in sheet[1]]
        sheet.delete_rows(2, sheet.max_row)
        clauses = ["1", "2", "3", "4", "4.1", "5", "5.1", "5.1.1", "5.1.2", "5.2", "5.2.1", "5.2.2", "5.3", "5.3.1", "5.3.2", "5.4", "5.4.1", "5.4.1.1", "6", "A", "A.1"]
        for index, clause in enumerate(clauses):
            values = {"Standard_No": "GB 38031-2025", "Standard_Name": "电动汽车用动力蓄电池安全要求",
                      "Draft_Version": "测试草案 2024-05-27", "Clause_No": clause,
                      "Parent_Clause_No": clause.rpartition(".")[0], "Level": clause.count(".") + 1,
                      "Sequence": index + 1, "Import_ID": f"test-{index}", "Title_CN": f"测试章节 {clause}",
                      "Content_CN": "测试正文" if index != 5 else "", "Clause_Type": "Annex" if clause == "A" else "Clause",
                      "Initial_Comment": "需确认交流电路绝缘要求" if clause == "5.2.1" else (f"文件批注 {clause}" if index in (2,3,4,12) else "")}
            sheet.append([values.get(header, "") for header in headers])
        output = io.BytesIO()
        book.save(output)
        book.close()
        self.template = output.getvalue()

    def tearDown(self):
        self.folder.cleanup()

    def upload(self, target=None, filename=TEMPLATE.name, content=None):
        body = {"file": (io.BytesIO(self.template if content is None else content), filename)}
        if target is not None:
            body["target"] = json.dumps(target)
        route = "/api/drafts/imports" if target is not None else "/api/drafts/imports/read"
        return self.client.post(route, data=body, content_type="multipart/form-data",
                                headers={"X-User": "Reviewer"})

    def target(self):
        return {"standard_id": 32, "version_name": "讨论稿", "sub_version_no": "test-1.0",
                "draft_date": "2024-05-27"}

    def snapshot(self, draft):
        return {str(a["id"]): {key: a[key] for key in ("content", "annotation_type", "updated_at")}
                for a in draft["annotations"] if a["chapter_id"] is not None}

    def test_lifecycle_export_includes_all_persisted_annotations_and_metadata(self):
        imported = self.upload(self.target()).json
        draft_id = imported["draft_id"]
        detail = self.client.get(f"/api/drafts/{draft_id}").json
        chapter = detail["chapters"][0]
        created = []
        for kind in ("Interpretation", "Proposed Change", "Question", "Recommendation", "Others"):
            response = self.client.post(f"/api/drafts/{draft_id}/annotations", json={
                "chapter_id": chapter["id"], "annotation_type": kind,
                "content": f"导出验证 {kind}\n第二行 | 多行内容"}, headers={"X-User": "ExportAuthor"})
            self.assertEqual(response.status_code, 201, response.json)
            created.append(response.json)
        whole = self.client.post(f"/api/drafts/{draft_id}/annotations", json={
            "annotation_type": "Recommendation", "content": "整份草案的导出建议"},
            headers={"X-User": "ExportAuthor"})
        self.assertEqual(whole.status_code, 201)
        edited = self.client.put(f"/api/drafts/{draft_id}/annotations/{created[0]['id']}", json={
            "annotation_type": "Interpretation", "content": "修改后的解释\n保留换行"},
            headers={"X-User": "ExportEditor"})
        self.assertEqual(edited.status_code, 200)
        # A new application/request reads the committed file for the export.
        reopened = create_app(self.db).test_client()
        response = reopened.get("/api/standards/32/lifecycle.md")
        self.assertEqual(response.status_code, 200)
        text = response.get_data(as_text=True)
        saved = reopened.get(f"/api/drafts/{draft_id}").json["annotations"]
        for annotation in saved:
            self.assertIn(f"批注 #{annotation['id']} · {annotation['annotation_type']}", text)
            for field in ("created_at", "created_by", "updated_at", "updated_by"):
                self.assertIn(annotation[field], text)
            for line in annotation["content"].splitlines():
                self.assertIn("> " + line, text)
        self.assertEqual(text.count("**批注 #"), len(saved))
        self.assertIn("整份草案批注（1 条）", text)
        self.assertIn("章节批注（5 条）", text)
        self.assertIn("###### 1 · 测试章节 1", text)
        self.assertNotIn("导入初始批注", text)
        self.assertNotIn("导出验证 Interpretation", text)
        # Another standard's annotations must not leak into this export.
        self.assertNotIn("逐机型复核电气图纸", text)

    def test_multiple_typed_annotations_persist_and_exchange(self):
        result = self.upload(self.target()).json
        draft_id = result["draft_id"]
        chapter = self.client.get(f"/api/drafts/{draft_id}").json["chapters"][0]
        ids = []
        for kind in ("Interpretation", "Proposed Change", "Question", "Recommendation", "Others"):
            response = self.client.post(f"/api/drafts/{draft_id}/annotations", json={
                "chapter_id": chapter["id"], "annotation_type": kind, "content": f"正文 {kind}"},
                headers={"X-User": "Reviewer"})
            self.assertEqual(response.status_code, 201, response.json)
            ids.append(response.json["id"])
        edited = self.client.put(f"/api/drafts/{draft_id}/annotations/{ids[0]}", json={
            "annotation_type": "Recommendation", "content": "修改后的建议"})
        self.assertEqual(edited.status_code, 200, edited.json)
        # Reopen the actual database file independently of Flask/request state.
        with Database(self.db) as conn:
            rows = [tuple(row) for row in conn.execute(
                "SELECT id,annotation_type,content FROM draft_annotation WHERE chapter_id=%s ORDER BY id",
                (chapter["id"],)).fetchall()]
        self.assertEqual(len(rows), 5)
        self.assertEqual(rows[0], (ids[0], "Recommendation", "修改后的建议"))
        reopened = create_app(self.db).test_client()
        self.assertEqual(len([a for a in reopened.get(f"/api/drafts/{draft_id}").json["annotations"]
                              if a["chapter_id"] == chapter["id"]]), 5)
        exported = self.client.get("/api/transfer/export?user=Reviewer")
        other_db = new_database()
        init_db(other_db, with_seed=True)
        other = create_app(other_db).test_client()
        imported = other.post("/api/transfer/import", data={
            "file": (io.BytesIO(exported.data), "backup.json"), "as_user": "Reviewer"})
        self.assertEqual(imported.status_code, 200, imported.json)
        transferred = other.get(f"/api/drafts/{draft_id}").json
        transferred_chapter = transferred["chapters"][0]
        actual = [(a["annotation_type"], a["content"]) for a in transferred["annotations"]
                  if a["chapter_id"] == transferred_chapter["id"]]
        self.assertCountEqual(actual, [(row[1], row[2]) for row in rows])

    def test_annotation_type_validation_in_api_and_required_in_database(self):
        draft_id = self.upload(self.target()).json["draft_id"]
        for invalid in (None, "", "Unknown", "comment"):
            body = {"content": "批注", "annotation_type": invalid}
            response = self.client.post(f"/api/drafts/{draft_id}/annotations", json=body)
            self.assertEqual(response.status_code, 400, response.json)
            self.assertEqual(response.json["field"], "annotation_type")
        valid = self.client.post(f"/api/drafts/{draft_id}/annotations", json={
            "content": "整体说明", "annotation_type": "Interpretation"})
        self.assertEqual(valid.status_code, 201, valid.json)
        response = self.client.put(f"/api/drafts/{draft_id}/annotations/{valid.json['id']}", json={
            "content": "不得保存", "annotation_type": "Unknown"})
        self.assertEqual(response.status_code, 400)
        with Database(self.db) as conn:
            with self.assertRaises(db.IntegrityError):
                conn.execute("UPDATE draft_annotation SET annotation_type=%s WHERE id=%s", (None, valid.json["id"]))

    def test_replacement_preserves_multiple_notes_and_detects_type_only_edit(self):
        first = self.upload(self.target()).json
        draft_id = first["draft_id"]
        detail = self.client.get(f"/api/drafts/{draft_id}").json
        chapter = detail["chapters"][0]
        ids = []
        for kind in ("Question", "Recommendation"):
            response = self.client.post(f"/api/drafts/{draft_id}/annotations", json={
                "chapter_id": chapter["id"], "annotation_type": kind, "content": kind})
            self.assertEqual(response.status_code, 201)
            ids.append(response.json["id"])
        detail = self.client.get(f"/api/drafts/{draft_id}").json
        target = {"standard_id": 32, "draft_id": draft_id, "replace_import_id": first["import_id"],
                  "keep_annotation_ids": ids, "annotation_snapshot": self.snapshot(detail), "comment_snapshot": self.client.get(f"/api/drafts/{draft_id}").json["linked_comments"]}
        edited = self.client.put(f"/api/drafts/{draft_id}/annotations/{ids[0]}", json={
            "content": "Question", "annotation_type": "Interpretation"})
        self.assertEqual(edited.status_code, 200)
        self.assertEqual(self.upload(target).status_code, 409)
        target["annotation_snapshot"] = self.snapshot(self.client.get(f"/api/drafts/{draft_id}").json)
        target["comment_snapshot"] = self.client.get(f"/api/drafts/{draft_id}").json["linked_comments"]
        response = self.upload(target)
        self.assertEqual(response.status_code, 201, response.json)
        current = self.client.get(f"/api/drafts/{draft_id}").json
        kept = [a for a in current["annotations"] if a["id"] in ids]
        self.assertEqual(len(kept), 2)
        self.assertEqual(len({a["chapter_id"] for a in kept}), 1)
        self.assertNotEqual(kept[0]["chapter_id"], chapter["id"])
        self.assertCountEqual([a["annotation_type"] for a in kept], ["Interpretation", "Recommendation"])

    def test_replacement_adds_incoming_note_once_and_keeps_other_notes(self):
        first = self.upload(self.target()).json
        draft_id = first["draft_id"]
        detail = self.client.get(f"/api/drafts/{draft_id}").json
        initial = next(a for a in detail["annotations"] if a["clause_no"] == "5.2.1")
        self.assertEqual(self.client.put(f"/api/drafts/{draft_id}/annotations/{initial['id']}", json={
            "annotation_type": "Interpretation", "content": "保留的解释"}).status_code, 200)
        replace_ids = []
        for kind in ("Question", "Recommendation"):
            response = self.client.post(f"/api/drafts/{draft_id}/annotations", json={
                "chapter_id": initial["chapter_id"], "annotation_type": kind, "content": kind})
            self.assertEqual(response.status_code, 201)
            replace_ids.append(response.json["id"])
        response = self.upload({"standard_id": 32, "draft_id": draft_id,
            "replace_import_id": first["import_id"], "replace_annotation_ids": replace_ids,
            "annotation_snapshot": self.snapshot(self.client.get(f"/api/drafts/{draft_id}").json),
            "comment_snapshot": self.client.get(f"/api/drafts/{draft_id}").json["linked_comments"]})
        self.assertEqual(response.status_code, 201, response.json)
        notes = [a for a in self.client.get(f"/api/drafts/{draft_id}").json["annotations"] if a["clause_no"] == "5.2.1"]
        self.assertEqual(len(notes), 2)
        self.assertIn(initial["id"], [a["id"] for a in notes])
        # 文件里的 Initial_Comment 类型未知，导入为 Others
        self.assertCountEqual([a["annotation_type"] for a in notes], ["Interpretation", "Others"])

    def test_template_read_and_current_chapters(self):
        source = self.upload()
        self.assertEqual(source.status_code, 200)
        self.assertEqual(source.json["chapter_count"], 21)
        self.assertEqual(source.json["standards"][0]["std_no"], "GB 38031-2025")
        result = self.upload(self.target())
        self.assertEqual(result.status_code, 201, result.json)
        draft = self.client.get(f"/api/drafts/{result.json['draft_id']}").json
        self.assertEqual(self.client.get("/api/standards/32").json["name_cn"], "电动汽车用动力蓄电池安全要求")
        self.assertEqual(draft["imports"][0]["source_version"], "测试草案 2024-05-27")
        self.assertEqual(len(draft["chapters"]), 21)
        self.assertEqual(draft["chapters"][10]["clause_no"], "5.2.1")
        excel_annotation = next(a for a in draft["annotations"] if a["clause_no"] == "5.2.1")
        self.assertEqual(excel_annotation["content"], "需确认交流电路绝缘要求")
        self.assertNotIn("origin", excel_annotation)
        self.assertEqual(draft["chapters"][17]["level"], 4)
        self.assertEqual(draft["chapters"][19]["clause_type"], "Annex")
        self.assertFalse(draft["chapters"][5]["content_cn"])
        self.assertEqual(self.client.get(f"/drafts/{draft['id']}").status_code, 200)

    def test_replacement_rebinds_selected_annotation_and_deletes_unselected(self):
        first = self.upload(self.target()).json
        draft_id = first["draft_id"]
        detail = self.client.get(f"/api/drafts/{draft_id}").json
        annotated_ids = {a["chapter_id"] for a in detail["annotations"] if a["chapter_id"] is not None}
        chapter = next(c for c in detail["chapters"] if c["id"] not in annotated_ids)
        annotation = self.client.post(f"/api/drafts/{draft_id}/annotations",
                                      json={"annotation_type": "Question", "chapter_id": chapter["id"], "content": "人工核对振动要求"})
        self.assertEqual(annotation.status_code, 201)
        annotation_id = annotation.json["id"]
        self.assertIn("created_at", annotation.json)
        self.assertIn("updated_at", annotation.json)
        self.assertEqual(self.client.post(f"/api/drafts/{draft_id}/annotations",
                                          json={"annotation_type": "Question", "chapter_id": chapter["id"], "content": "重复批注"}).status_code, 201)
        edited = self.client.put(f"/api/drafts/{draft_id}/annotations/{annotation_id}",
                                 json={"annotation_type": "Interpretation", "content": "已确认振动要求"})
        self.assertEqual(edited.status_code, 200)
        self.assertEqual(edited.json["content"], "已确认振动要求")
        blocked = self.upload({"standard_id": 32, "draft_id": draft_id})
        self.assertEqual(blocked.status_code, 409)
        self.assertEqual(len(self.client.get(f"/api/drafts/{draft_id}").json["imports"]), 1)
        second = self.upload({"standard_id": 32, "draft_id": draft_id,
                              "replace_import_id": first["import_id"],
                              "keep_annotation_ids": [annotation_id],
                              "annotation_snapshot": {str(a["id"]): {key: a[key] for key in ("content", "annotation_type", "updated_at")} for a in self.client.get(f"/api/drafts/{draft_id}").json["annotations"] if a["chapter_id"] is not None},
                              "comment_snapshot": self.client.get(f"/api/drafts/{draft_id}").json["linked_comments"]})
        self.assertEqual(second.status_code, 201, second.json)
        current = self.client.get(f"/api/drafts/{draft_id}").json
        self.assertEqual(len(current["chapters"]), 21)
        self.assertEqual(len(current["imports"]), 1)
        kept = next(a for a in current["annotations"] if a["content"] == "已确认振动要求")
        current_chapter = next(c for c in current["chapters"] if c["clause_no"] == chapter["clause_no"])
        self.assertEqual(kept["chapter_id"], current_chapter["id"])
        self.assertNotEqual(kept["chapter_id"], chapter["id"])
        self.assertEqual(kept["created_at"], annotation.json["created_at"])
        removed = self.upload({"standard_id": 32, "draft_id": draft_id,
                               "replace_import_id": second.json["import_id"],
                               "keep_annotation_ids": [],
                               "annotation_snapshot": {str(a["id"]): {key: a[key] for key in ("content", "annotation_type", "updated_at")} for a in current["annotations"] if a["chapter_id"] is not None},
                               "comment_snapshot": self.client.get(f"/api/drafts/{draft_id}").json["linked_comments"]})
        self.assertEqual(removed.status_code, 201, removed.json)
        self.assertFalse(any(a["content"] == "已确认振动要求"
                             for a in self.client.get(f"/api/drafts/{draft_id}").json["annotations"]))

    def test_invalid_file_and_different_standard_assignment(self):
        self.assertEqual(self.upload(filename="bad.txt", content=b"not xlsx").status_code, 400)
        before = self.client.get("/api/standards/1").json["name_cn"]
        assigned = self.upload({**self.target(), "standard_id": 1, "sub_version_no": "assigned-from-file"})
        self.assertEqual(assigned.status_code, 201, assigned.json)
        detail = self.client.get(f"/api/drafts/{assigned.json['draft_id']}").json
        self.assertEqual(detail["standard_id"], 1)
        self.assertEqual(detail["imports"][0]["source_standard_no"], "GB 38031-2025")
        self.assertEqual(self.client.get("/api/standards/1").json["name_cn"], before)
        self.assertEqual(self.client.get("/api/drafts/clauses/search").status_code, 404)

    def test_embedded_illustration_is_saved_served_and_exchanged(self):
        book = load_workbook(io.BytesIO(self.template))
        sheet = book["Import_Data"]
        sheet["K2"] = "图 1 试验装置"
        raw_image = io.BytesIO()
        PillowImage.new("RGB", (18, 18), "blue").save(raw_image, "PNG")
        raw_image.seek(0)
        sheet.add_image(Image(raw_image), "K2")
        excel = io.BytesIO()
        book.save(excel)
        book.close()
        result = self.upload(self.target(), content=excel.getvalue())
        self.assertEqual(result.status_code, 201, result.json)
        draft_id = result.json["draft_id"]
        detail = self.client.get(f"/api/drafts/{draft_id}").json
        self.assertEqual(detail["chapters"][0]["illustration"], "图 1 试验装置")
        self.assertEqual(len(detail["illustrations"]), 1)
        image = detail["illustrations"][0]
        # 附图保存在数据库之外的附件文件夹，浏览器按记录里的路径直接取文件
        self.assertTrue(image["url"].startswith("/files/drafts/"))
        import storage
        with self.client.application.app_context():
            stored_bytes=storage.read(image["url"].removeprefix("/files/"))
            self.assertTrue(stored_bytes.startswith(b"\x89PNG"))
        self.assertTrue(self.client.get(image["url"]).data.startswith(b"\x89PNG"))
        exported = self.client.get("/api/transfer/export?user=Reviewer")
        with tempfile.TemporaryDirectory() as folder:
            other_db = new_database()
            init_db(other_db, with_seed=True)
            other = create_app(other_db).test_client()
            imported = other.post("/api/transfer/import", data={
                "file": (io.BytesIO(exported.data), "backup.json"), "as_user": "Reviewer",
            }, content_type="multipart/form-data")
            self.assertEqual(imported.status_code, 200, imported.json)
            transferred = other.get(f"/api/drafts/{draft_id}").json
            self.assertEqual(len(transferred["illustrations"]), 1)
            copied = transferred["illustrations"][0]["url"]
            self.assertNotEqual(copied, image["url"])
            self.assertEqual(other.get(copied).data, stored_bytes)

    def test_changed_annotation_blocks_stale_import_and_removed_clause_deletes_formal_comment(self):
        first = self.upload(self.target()).json
        draft_id = first["draft_id"]
        formal = self.client.post("/api/comments", json={
            "standard_id": 32, "draft_id": draft_id, "clause_no": "5.1.2",
            "comment_text": "Formal feedback", "source_kind": "single", "comment_type": "Proposed Change", "status": "Draft",
            "submitted_by": "Reviewer",
        })
        self.assertEqual(formal.status_code, 201, formal.json)
        previous = self.client.post('/api/drafts', json={'standard_id': 32, 'version_name': '讨论稿',
            'sub_version_no': 'previous-test', 'draft_date': '2024-01-01'})
        self.assertEqual(previous.status_code, 201, previous.json)
        comparison = self.client.post('/api/comments', json={'standard_id': 32, 'source_kind': 'comparison',
            'comment_type': 'Question', 'comment_text': '请核对前后条款变化。', 'submitted_by': 'Reviewer',
            'before_draft_id': previous.json['id'], 'before_clause_no': '5.1.1',
            'after_draft_id': draft_id, 'after_clause_no': '5.1.2'})
        self.assertEqual(comparison.status_code, 201, comparison.json)
        book = load_workbook(io.BytesIO(self.template))
        sheet = book["Import_Data"]
        headers = {cell.value: cell.column for cell in sheet[1]}
        for row in list(sheet.iter_rows(min_row=2)):
            clause = row[headers["Clause_No"] - 1].value
            if clause == "5.2.1":
                row[headers["Initial_Comment"] - 1].value = "New imported note"
            elif clause == "5.1.2":
                sheet.delete_rows(row[0].row)
        output = io.BytesIO()
        book.save(output)
        content = output.getvalue()
        detail = self.client.get(f"/api/drafts/{draft_id}").json
        old = next(a for a in detail["annotations"] if a["clause_no"] == "5.2.1")
        target = {
            "standard_id": 32, "draft_id": draft_id, "replace_import_id": first["import_id"],
            "replace_annotation_ids": [old["id"]], "keep_annotation_ids": [],
            "annotation_snapshot": {str(a["id"]): {key: a[key] for key in ("content", "annotation_type", "updated_at")} for a in detail["annotations"] if a["chapter_id"] is not None},
            "comment_snapshot": detail["linked_comments"],
        }
        edited = self.client.put(f"/api/drafts/{draft_id}/annotations/{old['id']}",
                                 json={"annotation_type": "Proposed Change", "content": "A later edit"})
        self.assertEqual(edited.status_code, 200)
        stale = self.upload(target, content=content)
        self.assertEqual(stale.status_code, 409, stale.json)
        self.assertEqual(len(self.client.get(f"/api/drafts/{draft_id}").json["chapters"]), 21)
        target["annotation_snapshot"][str(old["id"])] = {key: edited.json[key] for key in ("content", "annotation_type", "updated_at")}
        target["comment_snapshot"] = self.client.get(f"/api/drafts/{draft_id}").json["linked_comments"]
        result = self.upload(target, content=content)
        self.assertEqual(result.status_code, 201, result.json)
        current = self.client.get(f"/api/drafts/{draft_id}").json
        self.assertEqual(len(current["chapters"]), 20)
        self.assertEqual(next(a for a in current["annotations"] if a["clause_no"] == "5.2.1")["content"], "New imported note")
        retained = self.client.get(f'/api/comments/{comparison.json["id"]}')
        self.assertEqual(retained.status_code, 200, retained.json)
        self.assertEqual(retained.json['after_clause_no'], '5.1.2')
        self.assertTrue(all(c["clause_no"] != "5.1.2" for c in current["linked_comments"]))
        self.assertEqual(self.client.get(f"/api/comments/{formal.json['id']}").status_code, 404)

    def test_data_exchange_keeps_existing_current_chapters(self):
        first = self.upload(self.target()).json
        draft_id = first["draft_id"]
        exported = self.client.get("/api/transfer/export?user=Reviewer")
        self.assertEqual(exported.status_code, 200)
        imported = self.client.post("/api/transfer/import", data={
            "file": (io.BytesIO(exported.data), "backup.json"), "as_user": "Reviewer",
        }, content_type="multipart/form-data")
        self.assertEqual(imported.status_code, 200, imported.json)
        current = self.client.get(f"/api/drafts/{draft_id}").json
        self.assertEqual(len(current["imports"]), 1)
        self.assertEqual(len(current["chapters"]), 21)
        self.assertEqual(len([a for a in current["annotations"] if a["chapter_id"] is not None]), 5)
        self.assertGreater(imported.json["skipped"], 0)
        with tempfile.TemporaryDirectory() as folder:
            other_db = new_database()
            init_db(other_db, with_seed=True)
            other = create_app(other_db).test_client()
            fresh_import = other.post("/api/transfer/import", data={
                "file": (io.BytesIO(exported.data), "backup.json"), "as_user": "Reviewer",
            }, content_type="multipart/form-data")
            self.assertEqual(fresh_import.status_code, 200, fresh_import.json)
            fresh_draft = other.get(f"/api/drafts/{draft_id}").json
            self.assertEqual(len(fresh_draft["imports"]), 1)
            self.assertEqual(len(fresh_draft["chapters"]), 21)
            self.assertEqual(len([a for a in fresh_draft["annotations"] if a["chapter_id"] is not None]), 5)

    def test_changed_formal_comment_blocks_stale_import(self):
        first = self.upload(self.target()).json
        draft_id = first["draft_id"]
        formal = self.client.post("/api/comments", json={
            "standard_id": 32, "draft_id": draft_id, "clause_no": "5.2.1",
            "comment_text": "Initial formal feedback", "source_kind": "single", "comment_type": "Proposed Change",
            "submitted_by": "Reviewer",
        })
        self.assertEqual(formal.status_code, 201, formal.json)
        detail = self.client.get(f"/api/drafts/{draft_id}").json
        target = {
            "standard_id": 32, "draft_id": draft_id, "replace_import_id": first["import_id"],
            "annotation_snapshot": {str(a["id"]): {key: a[key] for key in ("content", "annotation_type", "updated_at")} for a in detail["annotations"] if a["chapter_id"] is not None},
            "comment_snapshot": detail["linked_comments"],
        }
        changed = self.client.put(f"/api/comments/{formal.json['id']}", json={
            "comment_text": "Later formal feedback",
        })
        self.assertEqual(changed.status_code, 200, changed.json)
        stale = self.upload(target)
        self.assertEqual(stale.status_code, 409, stale.json)
        self.assertEqual(len(self.client.get(f"/api/drafts/{draft_id}").json["imports"]), 1)


    def test_annotation_delete_checks_creator_and_draft_scope(self):
        from urllib.parse import quote
        draft_id = self.upload(self.target()).json['draft_id']
        chapter = self.client.get(f'/api/drafts/{draft_id}').json['chapters'][0]
        for chapter_id in [None, chapter['id']]:
            created = self.client.post(f'/api/drafts/{draft_id}/annotations', json={
                'annotation_type': 'Proposed Change', 'content': '删除验证', 'chapter_id': chapter_id},
                headers={'X-User': quote('张三')}).json
            route = f'/api/drafts/{draft_id}/annotations/{created["id"]}'
            self.client.put(route, json={'annotation_type': 'Question', 'content': '他人编辑后仍归原作者'},
                            headers={'X-User': 'Editor'})
            self.assertEqual(self.client.delete(route, headers={'X-User': 'Editor'}).status_code, 403)
            self.assertEqual(self.client.delete(f'/api/drafts/1/annotations/{created["id"]}',
                             headers={'X-User': quote('张三')}).status_code, 404)
            self.assertEqual(self.client.delete(route, headers={'X-User': quote('张三')}).status_code, 200)
            self.assertEqual(self.client.delete(route, headers={'X-User': quote('张三')}).status_code, 404)
            self.assertNotIn(created['id'], [a['id'] for a in self.client.get(f'/api/drafts/{draft_id}').json['annotations']])

    def test_draft_group_size_is_capped_at_ten(self):
        for index in range(12):
            standard = self.client.post('/api/standards', json={
                'name_cn': f'分页标准{index}', 'stage_code': 'PRE_RESEARCH'}).json
            response = self.client.post('/api/drafts', json={'standard_id': standard['id'],
                'version_name': '讨论稿', 'sub_version_no': '1', 'draft_date': '2026-10-01'})
            self.assertEqual(response.status_code, 201, response.json)
        first = self.client.get('/api/drafts?page_size=50').json
        second = self.client.get('/api/drafts?page_size=50&page=2').json
        self.assertEqual(first['page_size'], 10)
        self.assertEqual(len(first['groups']), 10)
        self.assertLessEqual(len(second['groups']), 10)
        self.assertFalse({g['group_key'] for g in first['groups']} & {g['group_key'] for g in second['groups']})


if __name__ == "__main__":
    unittest.main()
