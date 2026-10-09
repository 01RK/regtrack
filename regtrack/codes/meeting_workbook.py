"""Read the three-sheet meeting Excel template at the upload boundary."""
from datetime import date, datetime
from io import BytesIO
from zipfile import ZipFile, BadZipFile
from xml.etree.ElementTree import ParseError
from openpyxl import load_workbook
from common import ApiError
from meeting_import import validate

SHEETS = {
    'Meeting': (('Meeting Title', 'Meeting Date', 'Key Discussions', 'Overall Conclusion', 'Organizer'),
                ('title', 'meeting_date', 'key_discussions', 'overall_conclusion', 'organizer')),
    'Standards': (('Standard Key', 'Standard Name', 'Standard No.', 'Minutes'),
                  ('key', 'name_cn', 'std_no', 'note')),
    'Actions': (('Action Title', 'Description', 'Standard Keys', 'Owner', 'Due Date'),
                ('title', 'description', 'standard_key', 'owner', 'due_date')),
}


def read_workbook(upload):
    if not upload or not upload.filename.lower().endswith('.xlsx'):
        raise ApiError('请上传 .xlsx 会议数据包，使用页面提供的 Excel 模板')
    content = upload.read(2 * 1024 * 1024 + 1)
    if len(content) > 2 * 1024 * 1024:
        raise ApiError('会议 Excel 文件不能超过 2 MB')
    try:
        with ZipFile(BytesIO(content)) as archive:
            if sum(item.file_size for item in archive.infolist()) > 20 * 1024 * 1024:
                raise ApiError('Excel 展开后过大，请只保留会议文本数据')
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=False, keep_links=False)
    except (BadZipFile, ValueError, KeyError, OSError, ParseError) as error:
        raise ApiError('无法读取 Excel，请使用模板另存为 .xlsx 后上传') from error
    try:
        if set(workbook.sheetnames) != set(SHEETS):
            raise ApiError('Excel 必须且只能包含 Meeting、Standards、Actions 三张工作表')
        result = {}
        for name, (headers, fields) in SHEETS.items():
            sheet = workbook[name]
            sheet.calculate_dimension(force=True)
            if sheet.max_row > 302 or sheet.max_column > len(headers):
                raise ApiError(f'{name} 超出模板范围：最多300条数据，请勿增加列')
            actual = tuple(c.value for c in next(sheet.iter_rows(min_row=1, max_row=1, max_col=len(headers))))
            if actual != headers:
                raise ApiError(f'{name} 第1行表头不符合模板，请保留原表头和列顺序')
            records = []
            for row in sheet.iter_rows(min_row=3, max_col=len(headers)):
                if all(c.value is None or c.value == '' for c in row):
                    continue
                record = {}
                for field, cell in zip(fields, row):
                    if cell.data_type in ('f', 'e'):
                        raise ApiError(f'{name}!{cell.coordinate} 含公式或错误值，请改为实际文本')
                    value = cell.value
                    if isinstance(value, (datetime, date)) and field in ('meeting_date', 'due_date'):
                        value = value.strftime('%Y-%m-%d')
                    if value is not None and not isinstance(value, str):
                        raise ApiError(f'{name}!{cell.coordinate} 请填写文本或日期，标准号请设为文本格式')
                    record[field] = value or ''
                records.append(record)
            result[name] = records
        if len(result['Meeting']) != 1:
            raise ApiError('Meeting 表第3行起必须且只能填写一场会议')
        return validate({'meeting': result['Meeting'][0], 'standards': result['Standards'], 'actions': result['Actions']})
    except (ValueError, KeyError, ParseError) as error:
        raise ApiError('Excel 内容损坏，请使用模板重新保存后上传') from error
    finally:
        workbook.close()
