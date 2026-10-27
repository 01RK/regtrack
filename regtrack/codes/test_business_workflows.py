from testdb import new_database
"""Current business and interface regression coverage."""
import io
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from app import create_app
from testdb import initialize as init_db
from constants import ANNOTATION_TYPES
from dbtest import Database

HEADERS = {"X-User": "Tester"}
TEMPLATE = Path(__file__).resolve().parents[1] / "static" / "standard-clause-import-blank.xlsx"


class BusinessWorkflowsTest(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.path = new_database()
        init_db(self.path, with_seed=True)
        self.client = create_app(self.path).test_client()

    # ------------------------------------------------------------ helpers
    def ok(self, response, status=200):
        self.assertEqual(response.status_code, status, response.get_json())
        return response.get_json()

    def collect_action(self, statuses=("Open", "Open")):
        recipients = [{"respondent_team": f"测试团队{i}", "response_status": s} for i, s in enumerate(statuses)]
        return self.ok(self.client.post("/api/actions", json={
            "item_type": "Collect Comments", "standard_id": 21, "title": "征集测试", "description": "征集测试",
            "recipients": recipients}, headers=HEADERS), 201)

    def set_status(self, recipient_id, value):
        return self.ok(self.client.put(f"/api/actions/recipients/{recipient_id}", json={
            "response_status": value, "response_actual_date": "2026-10-20"}, headers=HEADERS))

    def action(self, action_id):
        return self.ok(self.client.get(f"/api/actions/{action_id}"))

    # ------------------------------------------------- ⑥ 征集意见自动完成
    def test_collect_action_completes_when_no_recipient_is_open(self):
        action = self.collect_action()
        first, second = action["recipients"]
        self.set_status(first["id"], "Responded")
        self.assertNotEqual(self.action(action["id"])["current_status"], "Completed")
        self.set_status(second["id"], "No Comment")
        done = self.action(action["id"])
        self.assertEqual(done["current_status"], "Completed")
        self.assertIn("自动完成", done["status_history"][0]["note"])
        self.assertEqual(done["responded_count"], 2)
        # 又有人改回 Open：自动恢复为进行中
        self.set_status(second["id"], "Open")
        self.assertEqual(self.action(action["id"])["current_status"], "In Progress")

    def test_collect_action_created_with_all_replied_is_completed(self):
        action = self.collect_action(("Responded", "N-A"))
        self.assertEqual(action["current_status"], "Completed")

    def test_adding_open_recipient_reopens_completed_action(self):
        action = self.collect_action(("Responded",))
        self.ok(self.client.post(f"/api/actions/{action['id']}/recipients",
                                 json={"respondent_person": "新同事"}, headers=HEADERS), 201)
        self.assertEqual(self.action(action["id"])["current_status"], "In Progress")

    def test_manual_completion_blocked_while_recipients_open(self):
        action = self.collect_action()
        blocked = self.client.post(f"/api/actions/{action['id']}/status-history",
                                   json={"new_value": "Completed", "effective_date": "2026-10-20"}, headers=HEADERS)
        self.assertEqual(blocked.status_code, 400)
        self.assertIn("自动置为 Completed", blocked.get_json()["error"])
        self.ok(self.client.post(f"/api/actions/{action['id']}/status-history",
                                 json={"new_value": "Cancelled", "effective_date": "2026-10-20"}, headers=HEADERS), 201)

    def test_last_recipient_cannot_be_deleted(self):
        action = self.collect_action(("Open",))
        response = self.client.delete(f"/api/actions/recipients/{action['recipients'][0]['id']}", headers=HEADERS)
        self.assertEqual(response.status_code, 400)

    def test_no_opinion_generation_from_recipients(self):
        action = self.collect_action()
        recipient = action["recipients"][0]
        self.assertNotIn("generated_comment_no", recipient)
        self.assertEqual(self.client.get(f"/api/actions/recipients/{recipient['id']}/comment-draft").status_code, 404)
        with Database(self.path) as conn:
            self.assertNotIn("source_recipient_id", conn.columns("comment"))

    def test_seed_collect_actions_are_consistent(self):
        with Database(self.path) as conn:
            actions = conn.rows("SELECT id, current_status FROM action_item WHERE item_type = 'Collect Comments'")
            for action in actions:
                statuses = [r["response_status"] for r in conn.rows(
                    "SELECT response_status FROM feedback_recipient WHERE action_item_id = %s", (action["id"],))]
                self.assertTrue(statuses, action)
                self.assertEqual(action["current_status"] == "Completed", "Open" not in statuses, action)

    # ------------------------------------------------------- ③ 删除草案
    def test_draft_deletion_impact_and_delete(self):
        scope = self.ok(self.client.get("/api/drafts/15/deletion"))
        self.assertEqual((scope["chapters"], scope["images"], scope["comparisons"]), (2, 1, 0))
        self.assertGreater(scope["annotations"], 0)
        self.assertGreater(scope["comments"], 0)
        self.ok(self.client.delete("/api/drafts/15", headers=HEADERS))
        with Database(self.path) as conn:
            for table in ("draft_annotation", "draft_chapter", "draft_chapter_image"):
                self.assertEqual(conn.scalar(f"SELECT COUNT(*) FROM {table} WHERE draft_id = 15"), 0, table)
            self.assertEqual(conn.scalar("SELECT COUNT(*) FROM comment WHERE draft_id = 15"), 0)

    def test_draft_used_by_comparison_cannot_be_deleted(self):
        self.assertEqual(self.ok(self.client.get("/api/drafts/122/deletion"))["comparisons"], 1)
        response = self.client.delete("/api/drafts/122", headers=HEADERS)
        self.assertEqual(response.status_code, 409)
        self.assertIn("删除比对", response.get_json()["error"])

    # -------------------------------------------------- ⑧ 多段前言合并
    def workbook(self, rows):
        book = load_workbook(TEMPLATE)
        sheet = book["Import_Data"]
        headers = [cell.value for cell in sheet[1]]
        sheet.delete_rows(2, sheet.max_row)
        for index, row in enumerate(rows, 1):
            values = {"Standard_No": "ISO 1-1997", "Standard_Name": "Test", "Draft_Version": "未标明版本",
                      "Level": 1, "Sequence": index, "Import_ID": f"D-{index}", **row}
            sheet.append([values.get(header, "") for header in headers])
        out = io.BytesIO()
        book.save(out)
        return out.getvalue()

    def read(self, rows):
        return self.client.post("/api/drafts/imports/read",
                                data={"file": (io.BytesIO(self.workbook(rows)), "x.xlsx")})

    def test_adjacent_forewords_are_merged(self):
        result = self.read([
            {"Clause_No": "0", "Title_CN": "前言", "Title_EN": "Foreword", "Content_CN": "ISO 前言", "Source_Page": "ii"},
            {"Clause_No": "0", "Title_CN": "前言", "Title_EN": "Foreword", "Content_CN": "EN 前言", "Source_Page": "iv"},
            {"Clause_No": "0", "Title_CN": "引言", "Title_EN": "Introduction", "Content_CN": "引言"},
            {"Clause_No": "1", "Title_CN": "范围", "Content_CN": "范围"}])
        data = self.ok(result)
        self.assertEqual(data["chapter_nos"], ["前言", "引言", "1"])

    def test_real_duplicates_name_the_row(self):
        response = self.read([{"Clause_No": "1", "Title_CN": "范围"}, {"Clause_No": "1", "Title_CN": "范围二"}])
        self.assertEqual(response.status_code, 400)
        self.assertIn("第 3 行", response.get_json()["error"])
        response = self.read([{"Clause_No": "1", "Title_CN": "范围"}, {"Clause_No": "2", "Import_ID": "D-1"}])
        self.assertIn("Import_ID", response.get_json()["error"])

    # ------------------------------------------------------- ⑦ 批注类型
    def test_annotation_types(self):
        self.assertEqual(ANNOTATION_TYPES, ["Interpretation", "Proposed Change", "Question", "Recommendation", "Others"])
        for kind, status in (("Others", 201), ("Comment", 400)):
            response = self.client.post("/api/drafts/15/annotations", json={
                "annotation_type": kind, "content": "类型测试"}, headers=HEADERS)
            self.assertEqual(response.status_code, status, kind)
        with Database(self.path) as conn:
            self.assertEqual(conn.scalar("SELECT COUNT(*) FROM draft_annotation WHERE annotation_type = 'Comment'"), 0)
            self.assertEqual(conn.scalar("SELECT COUNT(*) FROM comment WHERE comment_type = 'Comment'"), 0)

    # ------------------------------------------------------- ⑤ 待办计数
    def test_standard_list_counts_todo_actions(self):
        items = {s["id"]: s for s in self.ok(self.client.get("/api/standards"))["items"]}
        with Database(self.path) as conn:
            expected = conn.scalar("""SELECT COUNT(*) FROM action_item WHERE standard_id = 21
                AND current_status NOT IN ('Completed', 'Closed', 'Cancelled')""")
        self.assertEqual(items[21]["todo_count"], expected)
        self.assertNotIn("open_action_count", items[21])

    # --------------------------------------------------- ① 操作人部门
    def test_new_operator_department_goes_into_team_dictionary(self):
        self.ok(self.client.post("/api/lookups/person", json={"value": "王学松", "note": "法规新部门"}, headers=HEADERS), 201)
        lookups = self.ok(self.client.get("/api/lookups"))
        self.assertIn("法规新部门", [row["value"] for row in lookups["team"]])
        person = next(row for row in lookups["person"] if row["value"] == "王学松")
        self.assertEqual(person["note"], "法规新部门")
        self.ok(self.client.post("/api/lookups/person", json={"value": "王学松", "note": "标准法规部"}, headers=HEADERS))
        person = next(row for row in self.ok(self.client.get("/api/lookups"))["person"] if row["value"] == "王学松")
        self.assertEqual(person["note"], "标准法规部")


if __name__ == "__main__":
    unittest.main()
