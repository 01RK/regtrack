"""Read the chapter template at the upload boundary; no database or file writes here."""
from collections import defaultdict
from zipfile import BadZipFile
from xml.etree.ElementTree import ParseError

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from common import ApiError

FIELDS = {
    "Clause_No": "clause_no", "Parent_Clause_No": "parent_clause_no",
    "Title_CN": "title_cn", "Content_CN": "content_cn", "Title_EN": "title_en",
    "Content_EN": "content_en", "Illustration": "illustration", "Clause_Type": "clause_type", "Source_Page": "source_page",
    "Initial_Comment": "initial_comment", "Review_Status": "review_status",
    "Confidence": "confidence", "Needs_Review": "needs_review", "Import_ID": "import_key",
}
REQUIRED = {"Standard_No", "Standard_Name", "Draft_Version", "Clause_No", "Level", "Sequence", "Import_ID"}


def text(value):
    return "" if value is None else str(value)


def merge_continuation(chapter, extra):
    """把相邻的同名前言 / 引言行并入前一行：正文、附图说明与批注依次接续，页码合并。"""
    for key in ("content_cn", "content_en", "illustration", "initial_comment"):
        parts = [part for part in (chapter[key].strip(), extra[key].strip()) if part]
        chapter[key] = "\n\n".join(parts)
    pages = [page for page in (chapter["source_page"].strip(), extra["source_page"].strip()) if page]
    chapter["source_page"] = ", ".join(dict.fromkeys(pages))


def special_clause_no(chapter):
    if chapter["level"] != 1:
        return None
    title_cn = chapter["title_cn"].strip()
    title_en = chapter["title_en"].strip().casefold()
    kind = chapter["clause_type"].strip().casefold()
    if title_cn == "前言":
        return "前言"
    if title_cn == "引言":
        return "引言"
    if title_en in {"foreword", "preface"} or kind in {"foreword", "preface"}:
        return "前言"
    if title_en == "introduction" or kind == "introduction":
        return "引言"
    return None


def read_upload(upload):
    if not upload or not upload.filename.lower().endswith(".xlsx"):
        raise ApiError("请选择章节模板的 .xlsx 文件")
    try:
        book = load_workbook(upload.stream, read_only=False, data_only=False, keep_links=False)
    except (BadZipFile, InvalidFileException, ValueError, KeyError, ParseError):
        raise ApiError("无法读取此 Excel，请选择有效的 .xlsx 文件") from None
    try:
        if "Import_Data" not in book.sheetnames:
            raise ApiError("文件缺少 Import_Data 工作表，请使用章节导入模板")
        sheet = book["Import_Data"]
        iterator = sheet.iter_rows()
        headers = [text(cell.value).strip() for cell in next(iterator, ())]
        if not REQUIRED.issubset(headers) or len([h for h in headers if h]) != len(set(h for h in headers if h)):
            raise ApiError("文件字段与章节模板不一致，请使用模板重新导出")
        result, source, seen = [], None, {key: set() for key in ("clause_no", "sequence", "import_key")}
        clause_by_row = {}
        raw_to_internal = defaultdict(set)
        for row in iterator:
            excel_row = row[0].row
            if any(cell.data_type == "f" for cell in row):
                raise ApiError("导入表含 Excel 计算公式；请粘贴为文本值。技术公式请用 $...$ 或 $$...$$ 的 LaTeX 文本表示")
            row = [cell.value for cell in row]
            if not any(v is not None and text(v).strip() for v in row):
                continue
            values = dict(zip(headers, row))
            if any(values.get(key) is None or not text(values[key]).strip() for key in REQUIRED):
                raise ApiError("章节必填信息不完整，请补全模板后重新上传")
            identity = (text(values["Standard_No"]).strip(), text(values["Standard_Name"]).strip(),
                        text(values["Draft_Version"]).strip())
            if source and identity != (source["standard_no"], source["standard_name"], source["source_version"]):
                raise ApiError("一个 Excel 只能包含同一标准的同一份草案")
            source = {"standard_no": identity[0], "standard_name": identity[1],
                      "source_version": identity[2]}
            chapter = {target: text(values.get(key)) for key, target in FIELDS.items()}
            for key in ("clause_no", "parent_clause_no", "import_key"):
                chapter[key] = chapter[key].strip()
            for key in ("Level", "Sequence"):
                try:
                    number = float(values[key])
                    if not number.is_integer() or number < 1:
                        raise ValueError
                    chapter[key.lower()] = int(number)
                except (TypeError, ValueError, OverflowError):
                    raise ApiError("章节层级和顺序必须是正整数") from None
            raw_clause_no = chapter["clause_no"]
            special = special_clause_no(chapter)
            chapter["clause_no"] = special or raw_clause_no
            raw_to_internal[raw_clause_no].add(chapter["clause_no"])
            for key, label in (("sequence", "Sequence"), ("import_key", "Import_ID")):
                if chapter[key] in seen[key]:
                    raise ApiError(f"第 {excel_row} 行的 {label}「{chapter[key]}」与前面的行重复，请修正后上传")
                seen[key].add(chapter[key])
            clause_by_row[excel_row] = chapter["clause_no"]
            # 原文有多段前言（如 ISO 前言与 EN 前言）时，大模型常拆成相邻的几行：合并为同一章。
            if special and result and result[-1]["clause_no"] == special:
                merge_continuation(result[-1], chapter)
                continue
            if chapter["clause_no"] in seen["clause_no"]:
                raise ApiError(f"第 {excel_row} 行的 Clause_No「{raw_clause_no}」与前面的行重复；"
                               "除相邻的多段前言 / 引言会自动合并外，章节编号不能重复")
            seen["clause_no"].add(chapter["clause_no"])
            result.append(chapter)
            if len(result) > 10000:
                raise ApiError("单份草案最多支持 10000 个章节项")
        if not result:
            raise ApiError("文件中没有章节内容")
        for raw_no, internal_nos in raw_to_internal.items():
            if len(internal_nos) > 1 and not (raw_no == "0" and internal_nos == {"前言", "引言"}):
                raise ApiError("文件存在重复的章节编号；仅前言和引言可以共用 0")
        for chapter in result:
            parent_no = chapter["parent_clause_no"]
            if parent_no:
                internal_nos = raw_to_internal.get(parent_no, set())
                if len(internal_nos) != 1:
                    raise ApiError("父章节编号不明确，请检查 Parent_Clause_No")
                chapter["parent_clause_no"] = next(iter(internal_nos))
        result.sort(key=lambda r: r["sequence"])
        previous = {}
        for chapter in result:
            parent = previous.get(chapter["parent_clause_no"])
            if ((chapter["parent_clause_no"] and (not parent or parent["level"] + 1 != chapter["level"]))
                    or (not chapter["parent_clause_no"] and chapter["level"] != 1)):
                raise ApiError("章节目录层级不完整，请确认父章节存在并排在子章节之前")
            previous[chapter["clause_no"]] = chapter
        illustrations = defaultdict(list)
        illustration_col = headers.index("Illustration") if "Illustration" in headers else None
        image_bytes = 0
        for image in sheet._images:
            anchor = image.anchor
            if illustration_col is None or not hasattr(anchor, "_from") or anchor._from.col != illustration_col:
                raise ApiError("请将附图插入 Import_Data 的 Illustration 列对应条款行")
            clause_no = clause_by_row.get(anchor._from.row + 1)
            if not clause_no:
                raise ApiError("Illustration 列中的图片没有对应的条款行")
            mime = {"png": "image/png", "jpeg": "image/jpeg", "gif": "image/gif"}.get(image.format)
            if not mime:
                raise ApiError("附图仅支持 PNG、JPEG 或 GIF 格式")
            data = image._data()
            image_bytes += len(data)
            if image_bytes > 40 * 1024 * 1024:
                raise ApiError("附图解压后总大小不能超过 40 MB")
            illustrations[clause_no].append({"mime_type": mime, "data": data})
        return source, result, illustrations
    finally:
        book.close()
