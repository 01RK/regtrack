from testdb import new_database
"""v21 comparison contract and matrix integration regression tests."""
import base64
import io
import json
import tempfile
import unittest
from pathlib import Path
from openpyxl import load_workbook
from app import create_app
from testdb import initialize as init_db
import db
from comparison_workbook import workbook
from comparison_service import pair
from dbtest import STAMP
import base64
import files

PNG = base64.b64decode('iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+aN1cAAAAASUVORK5CYII=')


class ComparisonTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.path = new_database()
        init_db(self.path, with_seed=False)
        self.app = create_app(self.path)
        self.client = self.app.test_client()
        with self.app.app_context():
            self.sid = db.insert('standard', dict(STAMP, std_no='GB/T TEST', name_cn='Test standard', stage_code='DRAFTING'))
            self.before = self.draft('1.0', '2026-01-01')
            self.after = self.draft('2.0', '2026-02-01')
            self.batch(self.before, ['1','2','3','4','5','6','7','8'])
            self.batch(self.after, ['1','2','3.1','3.2','4','5','6','7','9'])
            db.commit()
        self.rows = [
            ['R001','1','1','Unchanged','Unchanged','未变化','Unchanged','内容相同','Same content',''],
            ['R002','2','2','Modified','Substantive','阈值改变','Threshold changed','10改为20','10 to 20',''],
            ['R003','3','3.1;3.2','Split','Editorial','拆分','Split','要求不变','Requirements unchanged',''],
            ['R004','4;5','4','Merged','Editorial','合并','Merged','要求不变','Requirements unchanged',''],
            ['R005','6','5','Moved','Editorial','编号调整','Renumbered','要求不变','Requirements unchanged',''],
            ['R006','7','6','Modified','Uncertain','图示待核实','Figure needs review','图像不清晰','Image unclear','人工核实 / Review figure'],
            ['R007','8','','Deleted','Substantive','删除要求','Requirement deleted','取消约束','Constraint removed',''],
            ['R008','','7;9','Added','Substantive','增加要求','Requirements added','增加约束','Constraints added',''],
        ]

    def tearDown(self):
        self.tmp.cleanup()

    def draft(self, sub, date):
        return db.insert('draft', dict(STAMP, standard_id=self.sid, version_name='讨论稿', sub_version_no=sub, draft_date=date))

    def batch(self, did, nos):
        bid = db.insert('draft_import', dict(draft_id=did, filename='source.xlsx', source_version='test', source_standard_no='GB/T TEST', source_standard_name='Test', created_at='2026-01-01 09:00:00', created_by='tester'))
        for i, no in enumerate(nos, 1):
            ch = db.insert('draft_chapter', dict(draft_id=did, import_batch_id=bid, clause_no=no, parent_clause_no='', level=1, sequence=i, title_cn='测试', content_cn=f'要求 {no}', title_en='Test', content_en=f'Requirement {no}', import_key=no))
            if i == 1:
                db.insert('draft_chapter_image', dict(draft_id=did, chapter_id=ch, position=1, mime_type='image/png',
                                                      file_path=files.save(f'drafts/{did}', PNG, 'image/png')))
        return bid

    def file(self, rows=None, mutate=None):
        with self.app.app_context():
            b, a = pair(self.before, self.after)
            wb = load_workbook(workbook(b, a))
        for row in self.rows if rows is None else rows:
            wb['Changes'].append(row)
        if mutate: mutate(wb)
        data = io.BytesIO(); wb.save(data); data.seek(0)
        return data

    def upload(self, route='import', rows=None, mutate=None):
        return self.client.post('/api/comparisons/' + route, data=dict(before_draft_id=str(self.before), after_draft_id=str(self.after), file=(self.file(rows, mutate), 'result.xlsx')), headers={'X-User':'tester'})

    def imported(self):
        r = self.upload(); self.assertEqual(r.status_code, 201, r.json)
        return r.json['id']

    def test_end_to_end_export_import_and_sources(self):
        q = dict(before_draft_id=self.before, after_draft_id=self.after)
        r = self.client.get('/api/comparisons/pair', query_string=q)
        self.assertEqual(r.status_code, 200, r.json)
        self.assertEqual((r.json['before']['chapter_count'], r.json['after']['chapter_count'], r.json['comparison_id']), (8, 9, None))
        self.assertEqual(r.json['standard']['std_no'], 'GB/T TEST')
        r = self.client.get('/api/comparisons/workbook', query_string=q)
        self.assertEqual(r.status_code, 200)
        self.assertIn('.xlsx', r.headers['Content-Disposition'])
        wb = load_workbook(io.BytesIO(r.data))
        self.assertEqual(wb.sheetnames, ['Comparison', 'Changes', 'Before', 'After', 'Figures', 'Instructions'])
        self.assertEqual([c.value for c in wb['Before']['A']][1:], ['1','2','3','4','5','6','7','8'])
        self.assertEqual(wb['Before']['J2'].value, 'F001')
        self.assertEqual(wb['Figures'].max_row, 3)
        self.assertEqual(sum(len(ws._images) for ws in wb), 2)
        r = self.upload('preview'); self.assertEqual(r.status_code, 200, r.json)
        self.assertEqual(r.json['row_count'], 8)
        cid = self.imported()
        detail = self.client.get(f'/api/comparisons/{cid}').json
        self.assertEqual(len(detail['changes']), 8)
        self.assertEqual(detail['changes'][2]['after_nos'], ['3.1','3.2'])
        r = self.client.get(f'/api/comparisons/{cid}/export')
        self.assertEqual(r.status_code, 200)
        wb = load_workbook(io.BytesIO(r.data))
        self.assertEqual(wb['Changes'].max_row, 9)
        self.assertEqual(self.upload().status_code, 409)
        self.assertEqual(self.client.get('/api/comparisons/pair', query_string=q).json['comparison_id'], cid)

    def test_annotation_single_source_matrix_edit_delete(self):
        cid = self.imported()
        change = self.client.get(f'/api/comparisons/{cid}').json['changes'][2]
        r = self.client.post(f'/api/comparisons/changes/{change["id"]}/opinions', json=dict(content='Review split', annotation_type='Recommendation', needs_follow_up=True), headers={'X-User':'tester'})
        self.assertEqual(r.status_code, 201, r.json)
        oid = r.json['id']
        c = self.client.get(f'/api/comments/{oid}').json
        self.assertEqual(c['source_kind'], 'comparison')
        self.assertEqual(c['after_clause_no'], '3.1;3.2')
        self.assertEqual(c['status'], 'Submitted')
        r = self.client.put(f'/api/comments/{oid}', json={'comment_text':'Updated from matrix'})
        self.assertEqual(r.status_code, 200, r.json)
        d = self.client.get(f'/api/comparisons/{cid}').json
        self.assertEqual(d['changes'][2]['opinions'][0]['comment_text'], 'Updated from matrix')
        self.assertEqual(self.client.put(f'/api/comments/{oid}', json={'after_clause_no':'4'}).status_code, 400)
        self.assertEqual(self.client.delete(f'/api/comparisons/{cid}').status_code, 409)
        self.assertEqual(self.client.delete(f'/api/drafts/{self.before}').status_code, 409)
        self.assertEqual(self.client.delete(f'/api/comments/{oid}').status_code, 200)
        self.assertEqual(self.client.delete(f'/api/comparisons/{cid}').status_code, 200)

    def test_coverage_duplicates_and_bad_refs_atomic(self):
        cases = [self.rows[:-1], self.rows + [self.rows[0]],
                 [list(r) for r in self.rows]]
        cases[-1][1][1] = 'unknown'
        for rows in cases:
            r = self.upload(rows=rows)
            self.assertEqual(r.status_code, 400, r.json)
        with self.app.app_context():
            self.assertEqual(db.query_one('SELECT COUNT(*) AS n FROM draft_comparison')['n'], 0)

    def test_informational_metadata_and_formula(self):
        self.assertEqual(self.upload(mutate=lambda w: w.remove(w['Figures'])).status_code, 400)
        self.assertEqual(self.upload(mutate=lambda w: setattr(w['Comparison']['D2'], 'value', 999)).status_code, 201)
        self.assertEqual(self.upload(mutate=lambda w: setattr(w['Changes']['F2'], 'value', '=1+1')).status_code, 400)
        self.assertEqual(self.upload('preview', mutate=lambda w: setattr(w['Changes']['G2'], 'value', '')).status_code, 200)

    def test_workbook_ids_do_not_override_selected_target(self):
        def change_ids(wb):
            for cell in wb['Comparison'][2]:
                cell.value = 999
        response = self.upload(mutate=change_ids)
        self.assertEqual(response.status_code, 201, response.json)
        detail = self.client.get(f"/api/comparisons/{response.json['id']}").json
        self.assertEqual(detail['before_draft_id'], self.before)
        self.assertEqual(detail['after_draft_id'], self.after)
        with self.app.app_context():
            b, a = pair(self.before, self.after)
        self.assertEqual(detail['before_import_id'], b['import_id'])
        self.assertEqual(detail['after_import_id'], a['import_id'])

    def test_old_workbook_imports_after_source_reimport(self):
        data = self.file()
        with self.app.app_context():
            new_batch = self.batch(self.after, ['1','2','3.1','3.2','4','5','6','7','9'])
            db.commit()
        response = self.client.post('/api/comparisons/import', data=dict(
            before_draft_id=str(self.before), after_draft_id=str(self.after),
            file=(data, 'old-result.xlsx')))
        self.assertEqual(response.status_code, 201, response.json)
        detail = self.client.get(f"/api/comparisons/{response.json['id']}").json
        self.assertEqual(detail['after_import_id'], new_batch)

    def test_saved_comparison_validation_uses_saved_sources(self):
        cid = self.imported()
        with self.app.app_context():
            self.batch(self.after, ['new-source'])
            db.commit()
            from comparison_service import validate_stored
            validate_stored(cid)

    def test_pair_validation(self):
        for before, after in [(self.before,self.before),(self.after,self.before),(self.before,99999)]:
            r = self.client.get('/api/comparisons/workbook', query_string=dict(before_draft_id=before, after_draft_id=after))
            self.assertIn(r.status_code, (400,404), r.json)
        with self.app.app_context():
            empty = self.draft('3.0','2026-03-01'); db.commit()
        self.assertEqual(self.client.get('/api/comparisons/workbook', query_string=dict(before_draft_id=self.before,after_draft_id=empty)).status_code,400)

    def test_bad_shape_class_and_uncertain(self):
        for column,value in [(3,'Added'),(4,'Invalid')]:
            rows = [list(r) for r in self.rows]
            index = 5 if column == 9 else 1
            rows[index][column] = value
            r = self.upload(rows=rows)
            self.assertEqual(r.status_code,400,r.json)

    def test_many_to_many_and_false_unchanged(self):
        rows = [list(r) for r in self.rows]
        rows[3][2] = '4;5'; rows[3][3] = 'Reorganized'
        rows[4][2] = '7'; rows[7][2] = '9'
        self.assertEqual(self.upload(rows=rows).status_code,201)
        rows = [list(r) for r in self.rows]
        rows[1][3] = 'Unchanged'; rows[1][4] = 'Unchanged'
        with self.app.app_context():
            db.execute("UPDATE draft_chapter SET content_cn='different' WHERE draft_id=%s AND clause_no='2'", (self.after,)); db.commit()
        self.assertEqual(self.upload('preview', rows=rows).status_code,200)

    def test_transfer_roundtrip(self):
        cid = self.imported()
        change = self.client.get(f'/api/comparisons/{cid}').json['changes'][0]
        self.client.post(f'/api/comparisons/changes/{change["id"]}/opinions',json={'content':'transfer note','annotation_type':'Proposed Change'},headers={'X-User':'tester'})
        bundle = self.client.get('/api/transfer/export',query_string={'user':'tester'})
        self.assertEqual(bundle.status_code,200,bundle.json)
        packet=json.loads(bundle.data)
        self.assertEqual(len(packet['tables']['draft_comparison']),1)
        self.assertEqual(len(packet['tables']['comparison_change']),8)
        other_path=new_database(); init_db(other_path,with_seed=False)
        other=create_app(other_path).test_client()
        response=other.post('/api/transfer/import',data={'as_user':'receiver','file':(io.BytesIO(bundle.data),'data.json')})
        self.assertEqual(response.status_code,200,response.json)
        d=other.get('/api/comparisons/1').json
        self.assertEqual(len(d['changes']),8)
        self.assertEqual(d['changes'][0]['opinions'][0]['comment_text'],'transfer note')

    def test_source_replacement_and_date_are_protected(self):
        self.imported()
        template = Path(__file__).resolve().parents[1] / 'static' / 'standard-clause-import-blank.xlsx'
        wb = load_workbook(template); ws=wb['Import_Data']
        headers=[c.value for c in ws[1]]; ws.delete_rows(2,ws.max_row)
        row={'Standard_No':'GB/T TEST','Standard_Name':'Test standard','Draft_Version':'test',
             'Clause_No':'1','Parent_Clause_No':'','Level':1,'Sequence':1,'Import_ID':'new-1',
             'Title_CN':'测试','Content_CN':'替换内容','Clause_Type':'Clause'}
        ws.append([row.get(h,'') for h in headers]); data=io.BytesIO(); wb.save(data);data.seek(0)
        r=self.client.post('/api/drafts/imports',data={'file':(data,'source.xlsx'),'target':json.dumps({'draft_id':self.before,'standard_id':self.sid})})
        self.assertEqual(r.status_code,409,r.json)
        self.assertEqual(self.client.put(f'/api/drafts/{self.before}',json={'draft_date':'2026-01-02'}).status_code,409)
        with self.app.app_context():
            self.assertEqual(db.query_one('SELECT content_cn FROM draft_chapter WHERE draft_id=%s AND clause_no=%s',(self.before,'1'))['content_cn'],'要求 1')

    def test_manual_review_edit(self):
        cid = self.imported()
        self.assertEqual(self.client.get(f'/api/comparisons/{cid}/').status_code, 404)
        self.assertEqual(self.client.get(f'/drafts/comparisons/{cid}').status_code, 200)
        changes = {c['row_key']: c for c in self.client.get(f'/api/comparisons/{cid}').json['changes']}
        self.assertEqual((changes['R001']['identical'], changes['R005']['identical'], changes['R003']['identical']), (True, False, False))
        texts = dict(summary_cn='说明', summary_en='Summary', reason_cn='依据', reason_en='Reason')
        edit = lambda key, **body: self.client.put(f'/api/comparisons/changes/{changes[key]["id"]}', json=dict(texts, **body))
        # Review notes are user-controlled and retained unless explicitly changed.
        r = edit('R006', revision_class='Substantive', review_note='ignored')
        self.assertEqual(r.status_code, 200, r.json)
        self.assertEqual((r.json['revision_class'], r.json['review_note'], r.json['summary_cn']), ('Substantive', 'ignored', '说明'))
        self.assertEqual(edit('R006', revision_class='Uncertain', review_note='').status_code, 200)
        self.assertEqual(edit('R006', revision_class='Uncertain', review_note='核实图示').json['review_note'], '核实图示')
        # Analysis labels are user-controlled; editing class preserves the imported mapping type.
        for revision_class in ('Unchanged', 'Editorial'):
            response = edit('R002', revision_class=revision_class)
            self.assertEqual(response.status_code, 200, response.json)
            self.assertEqual(response.json['revision_class'], revision_class)
            # The edit endpoint returns analysis fields only. Read the persisted
            # mapping from the detail endpoint rather than assuming an old response shape.
            detail = self.client.get(f'/api/comparisons/{cid}').json
            saved = next(row for row in detail['changes'] if row['row_key'] == 'R002')
            self.assertEqual(saved['change_type'], 'Modified')
        self.assertEqual(edit('R005', revision_class='Unchanged').status_code, 200)
        self.assertEqual(edit('R003', revision_class='Unchanged').status_code, 200)
        self.assertEqual(edit('R004', revision_class='Unchanged').status_code, 200)
        self.assertEqual(edit('R004', revision_class='Editorial', summary_en='').status_code, 200)
        self.assertEqual(edit('R004', revision_class='Bogus').status_code, 400)
        with self.app.app_context():
            from comparison_service import validate_stored
            validate_stored(cid)

    def test_unchanged_compares_figure_count_not_figure_content(self):
        with self.app.app_context():
            image = db.query_one('SELECT * FROM draft_chapter_image WHERE draft_id=%s', (self.after,))
            import storage
            replacement=PNG[:-8]+b'\x00'*8
            storage.client().put_object(storage.bucket(),image['file_path'],io.BytesIO(replacement),len(replacement),content_type='image/png')
        cid = self.imported()
        changes = {c['row_key']: c for c in self.client.get(f'/api/comparisons/{cid}').json['changes']}
        self.assertTrue(changes['R001']['identical'])
        with self.app.app_context():
            db.insert('draft_chapter_image', dict(draft_id=self.after, chapter_id=image['chapter_id'], position=2, mime_type='image/png',
                                                  file_path=files.save(f'drafts/{self.after}', PNG, 'image/png')))
            db.commit()
        changes = {c['row_key']: c for c in self.client.get(f'/api/comparisons/{cid}').json['changes']}
        self.assertFalse(changes['R001']['identical'])

    def test_empty_and_corrupt_upload(self):
        self.assertEqual(self.upload(rows=[]).status_code,400)
        r = self.client.post('/api/comparisons/import',data=dict(before_draft_id=str(self.before),after_draft_id=str(self.after),file=(io.BytesIO(b'bad'), 'bad.xlsx')))
        self.assertEqual(r.status_code,400,r.json)


if __name__ == '__main__':
    unittest.main()
